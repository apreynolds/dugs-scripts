# requirements_io.py — load and validate a degree-requirements workbook.
# Usage: imported by transcript_to_tracker.py; not run directly.
#        (Requirements workbooks are maintained by hand; copy one of the
#        _TEMPLATE-*.xlsx files in degree-requirements/ to start a new one.)
#
# A degree-requirements file is the DATA behind a tracker: which categories a
# program has, how many credit hours each needs, and which courses satisfy
# them. Keeping it as a separate Excel file means:
#
#   - a colleague can maintain BA Math / Math Econ / Math Phys without
#     touching any code;
#   - a new calendar edition is a new file, not a new column bolted onto the
#     tracker sheet;
#   - the tracker itself becomes a generated artifact, so its formula ranges
#     are computed rather than hand-drawn (and therefore can't drift).
#
# The file is deliberately "dumb": plain values, no formulas. Every formula
# lives in the generated tracker.
#
# ---------------------------------------------------------------------------
# SHEET LAYOUT
# ---------------------------------------------------------------------------
#
# Sheet "Requirements" — one row per category, header row first:
#
#   Category | Description | Courses | <Track1> | <Track2> | ... | Note
#
#   Category     short unique code, e.g. "Calc"    (appears in the tracker's
#                Cat. dropdown)
#   Description  human text shown in the tracker,  e.g. "MATH 1003, MATH 1013"
#   Courses      MACHINE-READABLE list of satisfying courses, separated by
#                semicolons, e.g. "MATH 1003; MATH 1013". A "*" stands for one
#                digit, so "MATH 3***" means any 3000-level MATH course.
#                Leave blank when only an advisor can decide (Breadth, etc.).
#   <Track...>   one column per program track. The COLUMN HEADINGS ARE THE
#                TRACK NAMES — add a column to add a track. Values are the
#                required credit hours for that track.
#   Note         optional hover-tooltip text shown in the tracker.
#
#   Exactly one category is the "slack" category (named in Meta). Its required
#   hours are computed, not read, so the totals always reconcile.
#
# Sheet "Meta" — a two-column key/value list:
#
#   Program            BSc Mathematics
#   Edition            2020
#   Slack category     FreeElec
#   Total CH: Major    129
#   Total CH: Honours  132
#
# Sheet "Constraints" — rules a human must attest to, since they're unreliable
# to compute (e.g. whether 9 ch of Breadth form a coherent discipline):
#
#   # | Constraint | Tracks
#   1 | Max 6 ch STAT counted toward Breadth | (blank = applies to all)
#   2 | Min 3 ch 4000-level MATH             | Honours

import re
from pathlib import Path

import openpyxl

# Column headings in the Requirements sheet that are NOT track names.
RESERVED_HEADINGS = {"category", "description", "courses", "note", "notes"}

META_PROGRAM = "program"
META_EDITION = "edition"
META_SLACK = "slack category"
META_TOTAL_PREFIX = "total ch:"


class RequirementsError(Exception):
    """Raised when a requirements workbook is malformed.

    The message is shown directly to whoever ran the script, so it should name
    the file and say what to fix.
    """


class Category:
    """One requirement bucket, e.g. "9 ch of third-year MATH"."""

    def __init__(self, code, description, course_patterns, req_ch, note):
        self.code = code
        self.description = description
        self.course_patterns = course_patterns  # list of str, may be empty
        self.req_ch = req_ch                    # dict: track name -> int
        self.note = note

    def __repr__(self):
        return "<Category " + self.code + ">"


class Constraint:
    """A rule the advisor ticks off by hand rather than one we compute."""

    def __init__(self, number, text, tracks):
        self.number = number
        self.text = text
        self.tracks = tracks  # empty list = applies to every track

    def applies_to(self, track):
        return not self.tracks or track in self.tracks


class Requirements:
    """A whole degree-requirements file, loaded and validated."""

    def __init__(self, program, edition, slack_category, tracks,
                 categories, totals, constraints, source_path):
        self.program = program
        self.edition = edition
        self.slack_category = slack_category
        self.tracks = tracks              # list of track names, in file order
        self.categories = categories      # list of Category, in file order
        self.totals = totals              # dict: track name -> total CH
        self.constraints = constraints
        self.source_path = Path(source_path)

    @property
    def label(self):
        """Short name for the tracker's table heading, e.g. "2020"."""
        return str(self.edition)

    def category_codes(self):
        return [c.code for c in self.categories]

    def match_category(self, course_code):
        """Which category does `course_code` satisfy? Returns a code or None.

        The rule is **explicit matches beat wildcard matches**. Without that,
        MATH 3033 would match both "3rdMath" (listed explicitly) and
        "MathElec" (via the wildcard "MATH 3***"), and an ambiguity rule would
        blank out nearly every math course.

        Returns None when nothing matches, or when two categories match at the
        same level of specificity — those are left for the advisor rather than
        guessed at.
        """
        if not course_code:
            return None

        exact_hits = []
        wildcard_hits = []

        for category in self.categories:
            for pattern in category.course_patterns:
                if "*" in pattern:
                    if _wildcard_matches(pattern, course_code):
                        wildcard_hits.append(category.code)
                        break
                elif pattern.upper() == course_code.upper():
                    exact_hits.append(category.code)
                    break

        for hits in (exact_hits, wildcard_hits):
            unique = sorted(set(hits))
            if len(unique) == 1:
                return unique[0]
            if len(unique) > 1:
                return None  # genuinely ambiguous; advisor decides
        return None


def _wildcard_matches(pattern, course_code):
    """Match a course code against a pattern where "*" means one digit.

    "MATH 3***" matches "MATH 3033" but not "MATH 2033" or "CS 3033".
    """
    regex = "^" + re.escape(pattern).replace(r"\*", r"\d") + "$"
    return re.match(regex, course_code, flags=re.IGNORECASE) is not None


def _cell_text(value):
    """Normalise a cell to a trimmed string ("" for empty)."""
    if value is None:
        return ""
    return str(value).strip()


def _cell_int(value, where, path):
    """Read a cell that must hold a whole number of credit hours."""
    text = _cell_text(value)
    if text == "":
        return 0
    try:
        number = float(text)
    except ValueError:
        raise RequirementsError(
            str(path) + ": " + where + " should be a number, got " + repr(text)
        )
    if number != int(number):
        raise RequirementsError(
            str(path) + ": " + where + " should be a whole number of credit "
            "hours, got " + repr(text)
        )
    return int(number)


def _split_patterns(text):
    """Split the Courses cell into a list of course patterns.

    Accepts semicolons or commas as separators, since colleagues will
    reasonably type either.
    """
    if not text:
        return []
    parts = re.split(r"[;,]", text)
    return [p.strip().upper() for p in parts if p.strip()]


def _require_sheet(workbook, name, path):
    if name not in workbook.sheetnames:
        raise RequirementsError(
            str(path) + ": missing required sheet " + repr(name) +
            ". Found: " + ", ".join(workbook.sheetnames)
        )
    return workbook[name]


def _load_meta(sheet, path):
    """Read the Meta sheet's key/value pairs."""
    program = ""
    edition = ""
    slack = ""
    totals = {}

    for row in sheet.iter_rows(values_only=True):
        if not row:
            continue
        key = _cell_text(row[0]).lower()
        value = _cell_text(row[1]) if len(row) > 1 else ""
        if not key:
            continue

        if key == META_PROGRAM:
            program = value
        elif key == META_EDITION:
            edition = value
        elif key == META_SLACK:
            slack = value
        elif key.startswith(META_TOTAL_PREFIX):
            track = key[len(META_TOTAL_PREFIX):].strip()
            # Recover the track's original capitalisation from column A.
            original = _cell_text(row[0])[len(META_TOTAL_PREFIX):].strip()
            totals[original or track] = _cell_int(
                value, "Meta '" + _cell_text(row[0]) + "'", path
            )

    if not slack:
        raise RequirementsError(
            str(path) + ": Meta sheet must define 'Slack category' — the "
            "category that absorbs leftover credit hours (e.g. FreeElec)."
        )
    return program, edition, slack, totals


def _load_requirements_sheet(sheet, path):
    """Read the Requirements sheet into track names + Category objects."""
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise RequirementsError(str(path) + ": Requirements sheet is empty.")

    header = [_cell_text(cell) for cell in rows[0]]
    lowered = [h.lower() for h in header]

    for needed in ("category", "description", "courses"):
        if needed not in lowered:
            raise RequirementsError(
                str(path) + ": Requirements sheet needs a " + repr(needed) +
                " column. Found: " + ", ".join(h for h in header if h)
            )

    idx_category = lowered.index("category")
    idx_description = lowered.index("description")
    idx_courses = lowered.index("courses")
    idx_note = lowered.index("note") if "note" in lowered else None

    # Every remaining non-empty heading is a track name. This is what lets a
    # colleague add a "Co-op" track by adding one column.
    track_columns = {}
    for i, heading in enumerate(header):
        if heading and heading.lower() not in RESERVED_HEADINGS:
            track_columns[heading] = i

    if not track_columns:
        raise RequirementsError(
            str(path) + ": Requirements sheet has no track columns. Add at "
            "least one column (e.g. 'Major') holding required credit hours."
        )

    tracks = list(track_columns.keys())
    categories = []

    for row_number, row in enumerate(rows[1:], start=2):
        def cell(index):
            return _cell_text(row[index]) if index < len(row) else ""

        # Stop at the first entirely blank row. Anything below it is a
        # colleague's own notes, not data — reading on would turn a stray
        # comment into a phantom category.
        if not any(_cell_text(value) for value in row):
            break

        code = cell(idx_category)
        if not code:
            continue  # spacer row with content elsewhere; skip it

        req_ch = {}
        for track, column in track_columns.items():
            value = row[column] if column < len(row) else None
            req_ch[track] = _cell_int(
                value,
                "Requirements row " + str(row_number) + " track " +
                repr(track), path,
            )

        categories.append(Category(
            code=code,
            description=cell(idx_description),
            course_patterns=_split_patterns(cell(idx_courses)),
            req_ch=req_ch,
            note=cell(idx_note) if idx_note is not None else "",
        ))

    if not categories:
        raise RequirementsError(
            str(path) + ": Requirements sheet has no category rows."
        )
    return tracks, categories


def _load_constraints(workbook, path, tracks):
    """Read the optional Constraints sheet."""
    if "Constraints" not in workbook.sheetnames:
        return []

    sheet = workbook["Constraints"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    header = [_cell_text(c).lower() for c in rows[0]]
    idx_text = header.index("constraint") if "constraint" in header else 1
    idx_tracks = header.index("tracks") if "tracks" in header else None

    constraints = []
    for row in rows[1:]:
        if not row:
            continue
        text = _cell_text(row[idx_text]) if idx_text < len(row) else ""
        if not text:
            continue

        applies = []
        if idx_tracks is not None and idx_tracks < len(row):
            raw = _cell_text(row[idx_tracks])
            applies = [t.strip() for t in re.split(r"[;,]", raw) if t.strip()]

        for track in applies:
            if track not in tracks:
                raise RequirementsError(
                    str(path) + ": Constraints references unknown track " +
                    repr(track) + ". Known tracks: " + ", ".join(tracks)
                )

        constraints.append(
            Constraint(number=len(constraints) + 1, text=text, tracks=applies)
        )
    return constraints


def _validate(requirements, path):
    """Cross-checks that catch a colleague's edit before it reaches Excel."""
    codes = requirements.category_codes()

    duplicates = sorted({c for c in codes if codes.count(c) > 1})
    if duplicates:
        raise RequirementsError(
            str(path) + ": duplicate category codes: " + ", ".join(duplicates)
        )

    if requirements.slack_category not in codes:
        raise RequirementsError(
            str(path) + ": slack category " + repr(requirements.slack_category)
            + " is not one of the categories: " + ", ".join(codes)
        )

    for track in requirements.tracks:
        if track not in requirements.totals:
            raise RequirementsError(
                str(path) + ": Meta sheet is missing 'Total CH: " + track +
                "'. Every track column needs a total."
            )
        fixed = sum(
            c.req_ch.get(track, 0) for c in requirements.categories
            if c.code != requirements.slack_category
        )
        if fixed > requirements.totals[track]:
            raise RequirementsError(
                str(path) + ": track " + repr(track) + " requires " +
                str(fixed) + " ch across its categories, which exceeds the "
                "declared total of " + str(requirements.totals[track]) + " ch."
            )


def load_requirements(path):
    """Load and validate a requirements workbook. Returns a Requirements."""
    path = Path(path)
    if not path.exists():
        raise RequirementsError("No such requirements file: " + str(path))

    workbook = openpyxl.load_workbook(str(path), data_only=True)

    meta_sheet = _require_sheet(workbook, "Meta", path)
    req_sheet = _require_sheet(workbook, "Requirements", path)

    program, edition, slack, totals = _load_meta(meta_sheet, path)
    tracks, categories = _load_requirements_sheet(req_sheet, path)
    constraints = _load_constraints(workbook, path, tracks)

    requirements = Requirements(
        program=program,
        edition=edition or path.stem,
        slack_category=slack,
        tracks=tracks,
        categories=categories,
        totals=totals,
        constraints=constraints,
        source_path=path,
    )
    _validate(requirements, path)
    return requirements
