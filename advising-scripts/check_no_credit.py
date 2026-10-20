# /// script
# dependencies = ["openpyxl", "pdfplumber"]
# ///
# check_no_credit.py — test the strike-through on courses that earn nothing.
# Usage: uv run check_no_credit.py
#
# A regression check, not part of the shipped pipeline.
#
# The rule this covers: a course that does not earn its credit hours has its
# Course and Title cells greyed and struck through. F and W are the obvious
# cases; a D is the interesting one, because whether a D counts depends on the
# bar in that row's Pass cell, which depends on its category.
#
# The rule is therefore written as the NEGATION of the Progress test rather
# than as a list of bad grades — so the strike-through and the Progress column
# can never disagree about whether a course counted. That is the property
# worth protecting, and it is invisible in the generated XML: a hardcoded
# OR($E6="F",$E6="W") produces a well-formed workbook that looks right on
# every BSc and is wrong on the first BA row with a D in it.
#
# Three ways it can go wrong, none of them visible by inspection:
#
#   1. AN ARRAY CONSTANT. The compact MATCH($E6,{"D","C",...},0) idiom is the
#      natural way to rank a grade, and a conditional-formatting formula
#      cannot contain one. Excel does not complain about the rule — it
#      repairs the file and discards EVERY CF rule on the sheet, so the
#      requirements block silently loses its colours too. (Trap 1.)
#   2. AN ABSOLUTE ROW. A CF formula is written for the top-left cell of its
#      range and re-evaluated relatively. $E$6 instead of $E6 makes all 58
#      rows test row 6's grade, so the column strikes through as one.
#   3. THE WRONG SET OF ROWS. Striking out a blank or IP grade passes a
#      verdict on work that has not been marked; failing to strike an
#      "excl." row hides an exception the advisor made on purpose.
#
# What this canNOT check is whether Excel agrees. Nothing offline evaluates
# conditional formatting. See "Verifying a change" in CLAUDE.md — and note
# that check_formulas.py plants an F, a W, a bare D, a D + incl., an F + incl.
# and a B + incl. into its tracker, which is exactly the set of rows this rule
# has to get right. That workbook is the one to open.

import re
import subprocess
import sys
import zipfile
from pathlib import Path

import openpyxl

import check_expected_values as expected
import check_support
import tracker_writer

HERE = Path(__file__).resolve().parent

# What the rule MEANS, written out independently of the code that generates
# it: (grade, override, bar, struck through?).
#
# Spelled out rather than computed, for the same reason check_expected_values
# re-implements the grade tables by hand — a table derived from the code under
# test agrees with it however wrong it is.
INTENT = [
    ("F",  "",      "C", True,  "a failure earns nothing"),
    ("W",  "",      "C", True,  "nor does a withdrawal"),
    ("D",  "",      "C", True,  "below the bar for this category"),
    ("D",  "",      "D", False, "...but not below a bar of D"),
    ("D",  "incl.", "C", False, "rescued by hand: it counts, so it stands"),
    ("F",  "incl.", "C", True,  '"incl." must never resurrect an F'),
    ("C",  "",      "B", True,  "a C is struck where the bar is B"),
    ("A",  "excl.", "C", True,  "forced out — e.g. a superseded first try"),
    ("A",  "",      "C", False, "an ordinary pass"),
    ("CR", "",      "C", False, "transfer credit always counts"),
    ("",   "",      "C", False, "not taken yet — no verdict to pass"),
    ("IP", "",      "C", False, "in progress — likewise"),
]

# The seven things the generated formula is made of. Each entry is
# (fragment, what breaks if it goes missing). Fragments are checked against
# the formula as openpyxl reports it, i.e. already unescaped.
REQUIRED_FRAGMENTS = [
    ('$E6<>""',
     "a course with no grade yet would be struck through"),
    ('$E6<>"IP"',
     "a course in progress would be struck through"),
    ('($E6="A+")*9',
     "the grade has no rank, so nothing can be compared to the bar"),
    ('($F6="A+")*9',
     "the bar has no rank — the rule would be hardcoded to one programme"),
    ('$E6="CR"',
     "transfer credit would be struck through"),
    ('$G6<>"excl."',
     'an "excl." row would keep looking like it counts'),
    ('$G6="incl."',
     'an "incl." row would be struck through despite counting'),
]


def failing(message, failures):
    failures.append(message)
    print("  FAIL", message)


def build(work, name, args):
    """Generate a tracker and return its path, or None on failure."""
    tracker = work / name
    result = subprocess.run(
        [sys.executable, str(HERE / "transcript_to_tracker.py"),
         str(check_support.sample_transcript()), "-o", str(tracker)] + args,
        capture_output=True, text=True,
    )
    if result.returncode != 0:
        sys.stderr.write(result.stdout + result.stderr)
        return None
    return tracker


def find_rule(sheet):
    """The strike-through rule: (sqref, formula, dxf), or (None, None, None).

    Identified by its font — the only rule in the workbook that strikes text
    through without also filling the cell.
    """
    for rng in sheet.conditional_formatting:
        for rule in rng.rules:
            dxf = rule.dxf
            if dxf is not None and dxf.font is not None and dxf.font.strike \
                    and dxf.fill is None:
                return str(rng.sqref), rule.formula[0], dxf
    return None, None, None


def check_intent(failures):
    """The rule the sheet is supposed to implement, stated on its own."""
    print("1. what the rule means")

    for grade, override, bar, should_strike, why in INTENT:
        row = {"grade": grade, "ovr": override}
        # The sheet's condition, in Python: a recorded grade that earns
        # nothing. expected.counts_toward_progress is the independent oracle
        # check_formulas already measures the workbook against.
        struck = (grade not in ("", "IP")
                  and not expected.counts_toward_progress(row, bar))
        label = "{:2s} {:6s} bar {}".format(grade or "--", override or "", bar)
        if struck != should_strike:
            failing(label + " -> struck=" + str(struck) + ", expected " +
                    str(should_strike) + " (" + why + ")", failures)
        else:
            print("   OK  {}  ->  {:11s} {}".format(
                label, "struck out" if struck else "left alone", why))


def check_rule(tracker, kind, failures):
    """The generated rule, read back out of the workbook."""
    print()
    print(kind + ": the generated rule")

    workbook = openpyxl.load_workbook(str(tracker))
    sheet = workbook["Tracker"]
    sqref, formula, dxf = find_rule(sheet)

    if formula is None:
        failing("no strike-through rule in the " + kind + " tracker — "
                "failed and withdrawn courses look like passes", failures)
        return

    # --- the range ---------------------------------------------------------
    #
    # Resolved from the header row rather than assumed to be B:C, so the check
    # keeps working if a column is ever inserted ahead of Course.
    headings = {sheet.cell(row=5, column=c).value: c
                for c in range(1, sheet.max_column + 1)}
    course = openpyxl.utils.get_column_letter(headings["Course"])
    title = openpyxl.utils.get_column_letter(headings["Title"])
    last = sheet.max_row
    while last > 6 and sheet.cell(row=last, column=headings["gpa1"]).value \
            is None:
        last -= 1
    wanted = "$" + course + "$6:$" + title + "$" + str(last)
    if sqref.replace("$", "") != wanted.replace("$", ""):
        failing("the rule covers " + sqref + ", expected " + wanted +
                " — the Course and Title cells of every data row", failures)
    else:
        print("   OK  covers Course and Title on every data row (" +
              sqref + ")")

    # --- the three invisible faults ----------------------------------------

    if "{" in formula:
        failing("the formula contains an array constant — Excel will repair "
                "the file and discard EVERY conditional-formatting rule on "
                "the sheet: " + formula[:80], failures)
    else:
        print("   OK  no array constant")

    absolute = re.findall(r"\$[A-Z]+\$\d+", formula)
    if absolute:
        failing("absolute row references " + repr(sorted(set(absolute))) +
                " — every row would test row 6's grade", failures)
    else:
        print("   OK  references are column-absolute, row-relative")

    for fragment, breaks in REQUIRED_FRAGMENTS:
        if fragment not in formula:
            failing("the formula is missing " + fragment + " — " + breaks,
                    failures)
    if all(f in formula for f, _ in REQUIRED_FRAGMENTS):
        print("   OK  all " + str(len(REQUIRED_FRAGMENTS)) +
              " parts of the condition present")

    # The rescue term must stay bounded below the bar, exactly as it is in
    # the Progress SUMPRODUCT — same reason (check_formulas.py), and it keeps
    # the two statements of the rule textually parallel.
    rank = tracker_writer._grade_rank_cell("$E6")
    bar = tracker_writer._grade_rank_cell("$F6")
    if rank + "<" + bar not in formula:
        failing("the incl. rescue is not bounded below the bar — the strike-"
                "through would disagree with the Progress column", failures)
    else:
        print("   OK  the incl. rescue is bounded below the bar")

    # --- the format itself -------------------------------------------------

    colour = dxf.font.color.rgb if dxf.font.color is not None else None
    if not dxf.font.strike:
        failing("the format does not strike the text through", failures)
    elif colour != tracker_writer.NO_CREDIT_FONT.color.rgb:
        failing("the format's colour is " + repr(colour) + ", expected " +
                repr(tracker_writer.NO_CREDIT_FONT.color.rgb), failures)
    elif dxf.fill is not None:
        failing("the format carries a fill. Grey FILL means 'not on this "
                "track' and nothing else in this workbook; struck-out grey "
                "TEXT is what means 'does not count'", failures)
    else:
        print("   OK  grey struck-through text, no fill")


def check_xml(tracker, failures):
    """The same guard against the escaped XML, not just openpyxl's view."""
    print()
    print("3. the saved XML")
    with zipfile.ZipFile(tracker) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    formulas = re.findall(r"<formula>(.*?)</formula>", sheet_xml, re.S)
    if not formulas:
        failing("no conditional-formatting formulas survived the save",
                failures)
        return
    braces = [f for f in formulas if "{" in f]
    if braces:
        failing(str(len(braces)) + " CF formula(s) contain array constants",
                failures)
    else:
        print("   OK  " + str(len(formulas)) +
              " CF formulas saved, none with an array constant")


def main():
    failures = []
    work, pdf, req = check_support.prepare("_no_credit_test")

    check_intent(failures)

    tracker = build(work, "tracker.xlsx", ["-r", str(req), "-t", "Honours"])
    if tracker is None:
        failing("could not generate a tracker", failures)
    else:
        check_rule(tracker, "2. with requirements", failures)
        check_xml(tracker, failures)

    # The rule has to reach the ad hoc tracker too: that is the one where the
    # bar is typed by hand, so the strike-through is the only feedback that
    # a change to it did anything.
    ad_hoc = build(work, "adhoc.xlsx", ["--ad-hoc"])
    if ad_hoc is None:
        failing("could not generate an ad hoc tracker", failures)
    else:
        check_rule(ad_hoc, "4. ad hoc", failures)

    print()
    if failures:
        print("FAILED —", len(failures), "problem(s) above")
        return 1
    print("PASSED — courses that earn nothing are struck through, by the "
          "negation of the Progress test rather than a list of grades.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
