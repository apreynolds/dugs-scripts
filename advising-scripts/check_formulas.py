# /// script
# dependencies = ["openpyxl", "pdfplumber"]
# ///
# check_formulas.py — guard against array-unsafe formulas, and work an example.
# Usage: uv run check_formulas.py
#
# A regression check, not part of the shipped pipeline.
#
# SUMPRODUCT feeds a whole column into every grade test, so those tests must
# operate elementwise. Two idioms silently do NOT:
#
#   IFERROR(CHOOSE(MATCH(Courses[Gr.],{...},0),4.3,...),0)
#       CHOOSE rejects an array index -> #VALUE! -> IFERROR collapses the
#       whole factor to the SCALAR 0, so every GPA reads 0.00.
#
#   ISNUMBER(MATCH(Courses[Gr.],{...},0))
#       collapses to one TRUE/FALSE for the entire column, so Progress counts
#       every course in a category regardless of grade — failures included.
#
# Both were present in the first version of the generator and both are now
# written as sums of explicit (range="X") comparisons. This check fails if
# either idiom comes back, and then works a sub-GPA example end to end so the
# number can be compared against what Excel displays.

import re
import subprocess
import sys
import zipfile
from pathlib import Path

import openpyxl

import check_expected_values as expected
import check_support
import requirements_io
import tracker_reader

HERE = Path(__file__).resolve().parent

FORBIDDEN = {
    "CHOOSE(": "CHOOSE cannot take an array index; it collapses the GPA "
               "numerator to a scalar 0",
    "ISNUMBER(MATCH(": "MATCH does not broadcast here; Progress would count "
                       "failed courses",
}


def sheet_formulas(xlsx_path):
    """Pull every formula out of the workbook's XML, unescaped."""
    with zipfile.ZipFile(xlsx_path) as archive:
        names = [n for n in archive.namelist()
                 if n.startswith("xl/worksheets/sheet")]
        text = "".join(archive.read(n).decode("utf-8") for n in names)
    formulas = re.findall(r"<f>(.*?)</f>", text, re.S)
    unescape = {"&quot;": '"', "&gt;": ">", "&lt;": "<", "&amp;": "&"}
    out = []
    for formula in formulas:
        for old, new in unescape.items():
            formula = formula.replace(old, new)
        out.append(formula)
    return out


# (grade, override, should count toward Progress).
#
# The expectations are written out rather than computed, so this table is an
# independent statement of the rule. They hold because every category in the
# BSc fixture has a passing bar of C — the fixture carries no Pass column, so
# the default applies throughout. main() asserts that, rather than trusting it.
EDGE_CASES = [
    ("F", "", False),        # a failure earns nothing
    ("W", "", False),        # nor does a withdrawal
    ("D", "", False),        # below the bar for this programme
    ("D", "incl.", True),    # ...until the advisor allows it
    ("F", "incl.", False),   # but "incl." must never resurrect an F
    ("B", "incl.", True),    # already clears the bar: counted ONCE, not twice
]


def plant_edge_grades(sheet, headings):
    """Overwrite a few grades so failure/withdrawal semantics get exercised.

    Picks its targets by position rather than by course code, so it keeps
    working if the transcript fixture is ever replaced. Returns a list of
    (course, grade, override, should_count) describing what it changed.
    """
    wanted = list(EDGE_CASES)
    planted = []

    for row in range(6, 200):
        if not wanted:
            break
        course = sheet.cell(row=row, column=headings["Course"]).value
        grade = sheet.cell(row=row, column=headings["Gr."]).value
        term = sheet.cell(row=row, column=headings["Term"]).value
        # Only touch real graded rows. Skip transfer credits, where an "F"
        # would just look wrong to anyone reading the workbook, and skip the
        # MATH courses — they are the gpa1 sample below, which wants an
        # ordinary spread of grades.
        if not course or not grade or str(course).startswith("MATH"):
            continue
        if "TRANSFER" in str(term or "").upper():
            continue
        new_grade, override, should_count = wanted.pop(0)
        sheet.cell(row=row, column=headings["Gr."], value=new_grade)
        if override:
            sheet.cell(row=row, column=headings["Ovr."], value=override)
        planted.append((str(course), new_grade, override, should_count))

    if wanted:
        raise SystemExit(
            "the transcript fixture has too few graded non-MATH rows to plant "
            "edge-case grades in — " + str(len(wanted)) + " left over"
        )
    return planted


def main():
    WORK, PDF, REQ = check_support.prepare("_formula_test")
    tracker = WORK / "tracker.xlsx"

    result = subprocess.run(
        [sys.executable, str(HERE / "transcript_to_tracker.py"), str(PDF),
         "-r", str(REQ), "-o", str(tracker), "-t", "Honours"],
        capture_output=True, text=True,
    )
    if result.returncode != 0:
        print(result.stderr, file=sys.stderr)
        return 1

    print("1. array-safety guard")
    formulas = sheet_formulas(tracker)
    print("   " + str(len(formulas)) + " formulas in the workbook")
    failures = []
    for idiom, why in FORBIDDEN.items():
        hits = [f for f in formulas if idiom in f]
        if hits:
            failures.append(idiom + " is back (" + str(len(hits)) +
                            " formulas) — " + why)
        else:
            print("   OK  no " + idiom)

    # The "incl." rescue term must be bounded BELOW the passing bar. Without
    # that bound a course that already clears its bar satisfies both terms of
    # the Progress sum and SUMPRODUCT counts its credit hours TWICE — an error
    # no amount of reading the generated XML makes obvious, since the formula
    # stays perfectly well-formed. Nothing offline can evaluate these
    # formulas, so the guard is structural.
    rescues = [f for f in formulas if '(Courses[Ovr.]="incl.")' in f]
    unbounded = [f for f in rescues if ")<((Courses[Pass]=" not in f]
    if not rescues:
        failures.append('no formula rescues an "incl." row — the override '
                        "has stopped doing anything")
    elif unbounded:
        failures.append(
            "the incl. rescue is not bounded below the passing bar in " +
            str(len(unbounded)) + " formula(s) — a course that already "
            "passes would be counted twice")
    else:
        print("   OK  incl. rescue bounded below the bar (" +
              str(len(rescues)) + " formulas)")

    if failures:
        print()
        for failure in failures:
            print("  FAIL", failure)
        return 1

    # --- 2. work a sub-GPA example ------------------------------------------

    print()
    print("2. worked sub-GPA example")
    workbook = openpyxl.load_workbook(str(tracker))
    sheet = workbook["Tracker"]

    headings = {sheet.cell(row=5, column=c).value: c
                for c in range(1, sheet.max_column + 1)}
    gpa1_col = headings["gpa1"]

    # Plant the grades that make the manual Excel check DISCRIMINATING.
    #
    # sample-transcript.pdf is a good student: every grade is A-range or CR.
    # Against a transcript like that, a Progress column that wrongly counts
    # failed courses looks exactly like a correct one — there are no failures
    # to miscount — so trap 2 in CLAUDE.md would sail past the one step that
    # actually catches it, Patrick opening the workbook. Overwriting a few
    # grades here gives that step something to disagree about:
    #
    #   F            must not count toward Progress, and drags the GPA down
    #   W            must count toward neither (not even the GPA denominator)
    #   D            must not count toward Progress on its own...
    #   D + "incl."  ...but must once the advisor overrides it
    planted = plant_edge_grades(sheet, headings)

    flagged = []
    for row in range(6, 70):
        course = sheet.cell(row=row, column=headings["Course"]).value
        grade = sheet.cell(row=row, column=headings["Gr."]).value
        if course and str(course).startswith("MATH") and grade:
            sheet.cell(row=row, column=gpa1_col, value=True)
            flagged.append((course, grade,
                            sheet.cell(row=row,
                                       column=headings["CH"]).value))
    workbook.save(str(tracker))

    print("   planted edge-case grades:")
    for course, grade, override, _ in planted:
        print("     {:12s} {:3s} {}".format(course, grade, override or ""))
    print()
    print("   flagged " + str(len(flagged)) + " graded MATH courses as gpa1:")
    for course, grade, hours in flagged:
        print("     {:12s} {:3s} {} ch".format(course, str(grade), hours))

    data = tracker_reader.read_tracker(tracker)
    rows = data["rows"]

    print()
    print("   Expected values — compare these against Excel:")
    print("     D3  Overall GPA :", expected.gpa(rows))
    print("     I3  gpa1        :", expected.gpa(rows, "gpa1"),
          "  (MATH courses only)")
    print("     J3  gpa2        :", expected.gpa(rows, "gpa2"),
          "  (nothing flagged)")
    print()
    print("   Tracker saved at:", tracker)
    print()

    failures = []
    if expected.gpa(rows) == 0 or expected.gpa(rows, "gpa1") == 0:
        failures.append("a GPA computed as 0 — the points factor collapsed")

    # --- 3. the planted grades behave ---------------------------------------
    #
    # These assert the Python model in check_expected_values, which is what
    # Patrick's manual Excel comparison is measured against. If the model
    # itself stopped excluding failures, the comparison would agree with a
    # broken workbook.

    print("3. grade / override semantics")
    req = requirements_io.load_requirements(REQ)
    label = req.label
    by_course = {r["course"]: r for r in rows if r.get("course")}
    for course, grade, override, should_count in planted:
        row = by_course.get(course)
        if row is None:
            failures.append(course + " vanished from the tracker")
            continue
        # EDGE_CASES assumes a bar of C. Say so out loud rather than letting a
        # future fixture with a Pass column quietly invalidate the table.
        bar = expected.bar_for(row, req, label)
        if bar != "C":
            failures.append(
                course + " has a passing bar of " + bar + ", but EDGE_CASES "
                "is written for a bar of C")
            continue
        counts = expected.counts_toward_progress(row, bar)
        if counts != should_count:
            failures.append(
                "{} ({}{}) counts_toward_progress={} but should be {}".format(
                    course, grade, " + " + override if override else "",
                    counts, should_count))
        else:
            print("   OK  {:12s} {:3s}{:8s} -> Progress {}".format(
                course, grade, " " + override if override else "",
                "counts" if counts else "excluded"))

    withdrawn = [c for c, g, _, _ in planted if g == "W"]
    gpa_with = expected.gpa(rows)
    if any(by_course.get(c, {}).get("grade") == "W" for c in withdrawn):
        # A W must not reach the GPA denominator. Give it a huge credit-hour
        # value: if it were being counted, the average would move.
        inflated = [dict(r) for r in rows]
        for row in inflated:
            if row.get("course") in withdrawn:
                row["hours"] = 999
        if expected.gpa(inflated) != gpa_with:
            failures.append("a W grade is reaching the GPA denominator")
        else:
            print("   OK  W excluded from the GPA denominator")

    if failures:
        print()
        for failure in failures:
            print("  FAIL", failure)
        return 1
    print()
    print("PASSED — no array-unsafe idioms, GPAs non-zero, and failed / "
          "withdrawn courses excluded.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
