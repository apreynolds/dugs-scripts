# /// script
# dependencies = ["openpyxl", "pdfplumber"]
# ///
# check_adhoc.py — test the tracker built with no degree requirements at all.
# Usage: uv run check_adhoc.py
#
# A regression check, not part of the shipped pipeline.
#
# An ad hoc tracker is the Courses table on its own: somewhere to see what a
# student has taken, and flag / note / categorise it by hand, before there is
# a programme to measure it against. No requirements blocks, no track, no
# category dropdown, and a passing bar that is typed rather than derived.
#
# It is the same generator with an empty requirements list, which is the point
# — there is no second code path to keep in step — but that is also the risk:
#
#   1. EVERY blocks[0] IS A CRASH. The writer reaches into the first
#      requirements block for the Pass lookup, the track table and the
#      category dropdown. Each has to be guarded, and an unguarded one fails
#      loudly here rather than in front of a colleague.
#   2. THE CAT. COLUMN CHANGES NAME. "Cat." with no edition label, because
#      the categories are the advisor's own. tracker_reader keys on exactly
#      that difference, so the two have to agree or an ad hoc tracker's
#      categories vanish on its own -u re-run — silently, since a blank Cat.
#      cell is normal in a real tracker.
#   3. THE PASS COLUMN INVERTS. Everywhere else it is derived and must not be
#      typed in; here it is typed and nothing derives it. A leftover formula
#      would be #REF!, and a leftover "not for typing in" note would be
#      advice to ignore the only control the sheet has.
#
# Asking for it explicitly is part of the design: a missing -r is an error,
# not an ad hoc tracker. Same reasoning as the "—" sentinel in a requirements
# file — a forgotten flag must fail safe rather than quietly measuring the
# student against nothing.

import subprocess
import sys
from pathlib import Path

import openpyxl

import check_support
import tracker_reader
import tracker_writer

HERE = Path(__file__).resolve().parent


def failing(message, failures):
    failures.append(message)
    print("  FAIL", message)


def run(script, *args):
    return subprocess.run(
        [sys.executable, str(HERE / script)] + [str(a) for a in args],
        capture_output=True, text=True,
    )


def check_cli_guards(work, failures):
    """--ad-hoc has to be asked for, and cannot be combined with -r."""
    print("1. the flag has to be asked for")

    pdf = check_support.sample_transcript()
    req = work / "bsc-math-2020.xlsx"

    cases = [
        ("transcript_to_tracker.py", [pdf, "-o", work / "x.xlsx"],
         "--ad-hoc",
         "a forgotten -r must be an error, not a tracker with no targets"),
        ("transcript_to_tracker.py",
         [pdf, "--ad-hoc", "-r", req, "-o", work / "x.xlsx"],
         "cannot be combined",
         "--ad-hoc with -r is contradictory and must be refused"),
        ("blank_tracker.py", ["--ad-hoc"],
         "needs -o",
         "--ad-hoc has no requirements file to name the output after"),
        ("blank_tracker.py", [],
         "--ad-hoc",
         "blank_tracker with neither -r nor --ad-hoc must say so"),
    ]
    for script, args, wanted, why in cases:
        result = run(script, *args)
        if result.returncode == 0:
            failing(script + " " + " ".join(str(a) for a in args) +
                    " succeeded — " + why, failures)
        elif wanted not in result.stderr:
            failing(script + " refused it, but the message does not mention " +
                    repr(wanted) + ": " + result.stderr.strip()[-160:],
                    failures)
        else:
            print("   OK  " + why)


def check_structure(tracker, failures):
    """What is and is not in the generated workbook."""
    print()
    print("2. the generated workbook")

    workbook = openpyxl.load_workbook(str(tracker))

    if workbook.sheetnames != ["Tracker"]:
        failing("expected the Tracker sheet alone, got " +
                repr(workbook.sheetnames) + " — the hidden track lookup "
                "tables have nothing to hold", failures)
    else:
        print("   OK  one sheet, no hidden _TrackData")

    sheet = workbook["Tracker"]

    if sheet["A3"].value or sheet["B3"].value:
        failing("row 3 still carries a Track selector (" +
                repr(sheet["A3"].value) + ", " + repr(sheet["B3"].value) +
                ") — with no requirements there are no tracks, and a number "
                "in B3 would read back as a track name on the next merge",
                failures)
    else:
        print("   OK  no track label or selector")

    headings = {sheet.cell(row=5, column=c).value: c
                for c in range(1, sheet.max_column + 1)}

    labelled = [h for h in headings if h and h.startswith(
        tracker_reader.CAT_PREFIX)]
    if labelled:
        failing("found edition-labelled category columns " + repr(labelled) +
                " in a tracker with no editions", failures)
    elif tracker_writer.AD_HOC_CAT_COLUMN not in headings:
        failing("no " + repr(tracker_writer.AD_HOC_CAT_COLUMN) +
                " column — there is nowhere to categorise a course",
                failures)
    else:
        print("   OK  a single unlabelled " +
              repr(tracker_writer.AD_HOC_CAT_COLUMN) + " column")

    for row in range(1, 5):
        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(row=row, column=column).value
            if str(value or "").startswith("DEGREE REQUIREMENTS"):
                failing("a requirements block was written anyway, at " +
                        sheet.cell(row=row, column=column).coordinate,
                        failures)

    # --- the Pass column, which is the one that inverts --------------------

    pass_col = headings[tracker_writer.PASS_COLUMN]
    values = [sheet.cell(row=r, column=pass_col).value
              for r in range(6, min(sheet.max_row, 30) + 1)]
    formulas = [v for v in values if str(v or "").startswith("=")]
    if formulas:
        failing(str(len(formulas)) + " Pass cells still hold a formula — "
                "there is no category table for it to look into", failures)
    elif set(values) != {tracker_writer.DEFAULT_PASS_GRADE}:
        failing("Pass cells hold " + repr(sorted(set(map(str, values)))) +
                ", expected every one to start at " +
                tracker_writer.DEFAULT_PASS_GRADE, failures)
    else:
        print("   OK  every Pass cell is a plain, editable " +
              tracker_writer.DEFAULT_PASS_GRADE)

    # --- validations: the Cat. column loses one, the Pass column gains one --

    cat_letter = openpyxl.utils.get_column_letter(
        headings[tracker_writer.AD_HOC_CAT_COLUMN])
    pass_letter = openpyxl.utils.get_column_letter(pass_col)
    ranges = {}
    for validation in sheet.data_validations.dataValidation:
        for cells in str(validation.sqref).split():
            ranges[cells.split(":")[0].rstrip("0123456789")] = \
                validation.formula1

    if cat_letter in ranges:
        failing("the ad hoc Cat. column has a dropdown (" +
                str(ranges[cat_letter]) + ") — there is no list of "
                "categories to offer, so it can only get in the way",
                failures)
    else:
        print("   OK  no dropdown on the ad hoc Cat. column")

    if pass_letter not in ranges:
        failing("the Pass column has no dropdown, but it is now typed by "
                "hand — a mistyped bar would silently rank 0 and wave every "
                "grade through", failures)
    elif "F" in str(ranges[pass_letter]).replace('"', "").split(","):
        failing("the Pass dropdown offers F, which ranks below every grade — "
                "a bar of F waves through blanks, withdrawals and failures",
                failures)
    else:
        print("   OK  the Pass column has a grade dropdown, without F")

    if "B3" in ranges.get("B", "") or any(
            "B3" in str(v.sqref) for v in sheet.data_validations.dataValidation):
        failing("a track dropdown was attached to B3 anyway", failures)
    else:
        print("   OK  no track dropdown")

    # --- the one number it can still report --------------------------------

    earned = str(sheet["D2"].value or "")
    if "SUMPRODUCT" not in earned or "Courses[CH]" not in earned:
        failing("D2 does not hold the CH-earned total: " + repr(earned),
                failures)
    elif sheet["C2"].value != "CH earned":
        failing("D2 has no label beside it in C2", failures)
    else:
        print("   OK  a CH earned total, the only total with no programme "
              "behind it")


def check_round_trip(work, tracker, failures):
    """The advisor's own categories survive an -u re-run."""
    print()
    print("3. round trip")

    workbook = openpyxl.load_workbook(str(tracker))
    sheet = workbook["Tracker"]
    headings = {sheet.cell(row=5, column=c).value: c
                for c in range(1, sheet.max_column + 1)}
    cat_col = headings[tracker_writer.AD_HOC_CAT_COLUMN]

    # Plant the advisor's work on the first three real rows.
    planted = {}
    row = 6
    while len(planted) < 3 and row < 60:
        course = sheet.cell(row=row, column=headings["Course"]).value
        if course:
            sheet.cell(row=row, column=cat_col, value="my-category")
            sheet.cell(row=row, column=headings["Notes"],
                       value="hand-written note")
            sheet.cell(row=row, column=headings["Ovr."], value="incl.")
            sheet.cell(row=row, column=headings["gpa1"], value=True)
            planted[str(course)] = True
        row += 1
    workbook.save(str(tracker))

    if len(planted) != 3:
        failing("could not find three graded rows to plant on", failures)
        return

    data = tracker_reader.read_tracker(tracker)
    keyed = {r["course"]: r for r in data["rows"] if r.get("course")}
    missing = [c for c in planted
               if keyed.get(c, {}).get(tracker_reader.AD_HOC_CAT_KEY)
               != "my-category"]
    if missing:
        failing("tracker_reader did not read the ad hoc Cat. column back "
                "for " + repr(missing) + " — it reads 'Cat. <edition>' "
                "columns and would have skipped the bare one", failures)
    else:
        print("   OK  the bare Cat. column reads back under " +
              repr(tracker_reader.AD_HOC_CAT_KEY))

    result = run("transcript_to_tracker.py", check_support.sample_transcript(),
                 "--ad-hoc", "-u", tracker)
    if result.returncode != 0:
        sys.stderr.write(result.stdout + result.stderr)
        failing("re-running with -u --ad-hoc failed", failures)
        return

    data = tracker_reader.read_tracker(tracker)
    keyed = {r["course"]: r for r in data["rows"] if r.get("course")}
    lost = []
    for course in planted:
        row = keyed.get(course, {})
        if row.get(tracker_reader.AD_HOC_CAT_KEY) != "my-category":
            lost.append(course + " (Cat.)")
        if row.get("notes") != "hand-written note":
            lost.append(course + " (Notes)")
        if row.get("ovr") != "incl.":
            lost.append(course + " (Ovr.)")
        if not row.get("gpa1"):
            lost.append(course + " (gpa1)")
    if lost:
        failing("the merge lost " + ", ".join(lost), failures)
    else:
        print("   OK  Cat., Notes, Ovr. and gpa1 all survive an -u re-run")


def check_no_leak(work, failures):
    """A tracker WITH requirements is untouched by any of this."""
    print()
    print("4. the ordinary tracker is unchanged")

    tracker = work / "with-req.xlsx"
    result = run("transcript_to_tracker.py", check_support.sample_transcript(),
                 "-r", work / "bsc-math-2020.xlsx", "-o", tracker,
                 "-t", "Honours")
    if result.returncode != 0:
        sys.stderr.write(result.stdout + result.stderr)
        failing("could not generate a tracker with requirements", failures)
        return

    workbook = openpyxl.load_workbook(str(tracker))
    sheet = workbook["Tracker"]
    headings = [sheet.cell(row=5, column=c).value
                for c in range(1, sheet.max_column + 1)]

    if tracker_writer.AD_HOC_CAT_COLUMN in headings:
        failing("the bare ad hoc Cat. column leaked into a tracker that has "
                "requirements: " + repr(headings), failures)
    else:
        print("   OK  its category columns are still edition-labelled")

    if "_TrackData" not in workbook.sheetnames:
        failing("the hidden track lookup sheet went missing", failures)
    elif sheet["A3"].value != "Track" or not sheet["B3"].value:
        failing("the track selector went missing from B3", failures)
    elif sheet["C2"].value == "CH earned":
        failing("the ad hoc CH-earned readout leaked into a tracker that has "
                "a TOTAL row to report it", failures)
    else:
        print("   OK  track selector, hidden lookup sheet, no stray readout")


def check_blank_tracker(work, failures):
    """blank_tracker --ad-hoc gets the same workbook with no rows."""
    print()
    print("5. blank_tracker --ad-hoc")

    tracker = work / "blank-adhoc.xlsx"
    result = run("blank_tracker.py", "--ad-hoc", "-o", tracker,
                 "--rows", "12", "-n", "Last, First", "--id", "1234567")
    if result.returncode != 0:
        sys.stderr.write(result.stdout + result.stderr)
        failing("blank_tracker --ad-hoc failed", failures)
        return

    workbook = openpyxl.load_workbook(str(tracker))
    sheet = workbook["Tracker"]
    headings = [sheet.cell(row=5, column=c).value
                for c in range(1, sheet.max_column + 1)]
    if tracker_writer.AD_HOC_CAT_COLUMN not in headings:
        failing("no ad hoc Cat. column in the blank tracker", failures)
    elif sheet["B1"].value != "Last, First":
        failing("the student name did not reach B1", failures)
    else:
        print("   OK  same workbook, no course rows")


def main():
    failures = []
    work, pdf, req = check_support.prepare("_adhoc_test")

    check_cli_guards(work, failures)

    tracker = work / "adhoc.xlsx"
    result = run("transcript_to_tracker.py", pdf, "--ad-hoc", "-o", tracker)
    if result.returncode != 0:
        sys.stderr.write(result.stdout + result.stderr)
        failing("could not build an ad hoc tracker at all", failures)
    else:
        check_structure(tracker, failures)
        check_round_trip(work, tracker, failures)

    check_no_leak(work, failures)
    check_blank_tracker(work, failures)

    print()
    if failures:
        print("FAILED —", len(failures), "problem(s) above")
        return 1
    print("PASSED — an ad hoc tracker is the Courses table alone, its "
          "categories survive a re-run, and nothing leaks either way.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
