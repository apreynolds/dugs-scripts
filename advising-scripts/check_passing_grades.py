# /// script
# dependencies = ["openpyxl", "pdfplumber"]
# ///
# check_passing_grades.py — test per-category passing grades.
# Usage: uv run check_passing_grades.py
#
# A regression check, not part of the shipped pipeline.
#
# The rule this covers: a D earns credit in some categories and not in others,
# and which is which is DATA, set per category in the requirements file.
#
#   BSc   refuses a D everywhere. Its fixture carries no Pass column at all,
#         so every bar falls back to the default C — byte for byte the
#         behaviour every tracker had before this feature existed.
#   BA    accepts a D outside the major. Its fixture sets the default to D and
#         overrides the eight MATH/STAT categories back to C.
#
# Between them the two fixtures cover both paths, which is the reason the BA
# fixture exists.
#
# What this canNOT check is whether Excel agrees: nothing offline evaluates
# structured table references (Courses[Pass]). So this checks the inputs to
# the formula — the bars written into the sheet, and the shape of the lookup —
# and the Python model in check_expected_values that Patrick's manual
# comparison is measured against. See "Verifying a change" in CLAUDE.md.

import subprocess
import sys
from pathlib import Path

import openpyxl

import check_expected_values as expected
import check_support
import requirements_io
import tracker_reader

HERE = Path(__file__).resolve().parent

# Category -> the bar it should resolve to, for the BA fixture. Written out
# rather than read back from the file, so this states the intent independently.
BA_EXPECTED_BARS = {
    "ARTS": "D", "Grp1": "D", "Grp2": "D", "CS": "D", "FreeElec": "D",
    "Calc": "C", "2ndMath": "C", "3rdMath": "C", "Analys": "C", "Stat": "C",
    "MathElec": "C", "AnalysH": "C", "Thesis": "C",
}

# (grade, override, bar, should count). The pair that matters is the two D
# rows: the SAME grade earns credit under a D bar and earns nothing under a C
# bar. That is the whole feature in two lines.
SEMANTICS = [
    ("D", "", "D", True),       # a D in an elective, on a BA
    ("D", "", "C", False),      # the same D in the major
    ("D", "incl.", "C", True),  # ...unless the advisor allows it
    ("F", "incl.", "D", False),  # but an F is never rescued, at any bar
    ("F", "", "D", False),
    ("", "incl.", "D", False),  # nor is a planned course with no grade yet
    ("W", "incl.", "D", False),
    ("CR", "", "C", True),      # transfer credit always counts
    ("CR", "excl.", "C", False),
    ("A", "incl.", "C", True),  # already clears the bar: counted, once
    ("A", "excl.", "C", False),
    ("C", "", "C", True),       # exactly at the bar
]


def failing(message, failures):
    failures.append(message)
    print("  FAIL", message)


def check_schema(failures):
    """The two fixtures load with the bars their programmes imply."""
    print("1. requirements schema")
    work = check_support.fresh_dir("_passing_test")

    bsc = requirements_io.load_requirements(check_support.sample_requirements(work))
    bars = {c.code: bsc.pass_grade_for(c.code) for c in bsc.categories}
    if bsc.default_pass_grade != "C":
        failing("BSc fixture default is " + bsc.default_pass_grade +
                ", expected C (it carries no Default passing grade row)",
                failures)
    elif set(bars.values()) != {"C"}:
        failing("BSc fixture should have a bar of C everywhere, got " +
                repr(sorted(set(bars.values()))), failures)
    else:
        print("   OK  BSc: no Pass column -> every bar is the default C")

    ba = requirements_io.load_requirements(check_support.ba_requirements(work))
    if ba.default_pass_grade != "D":
        failing("BA fixture default is " + ba.default_pass_grade +
                ", expected D", failures)
    ba_bars = {c.code: ba.pass_grade_for(c.code) for c in ba.categories}
    if ba_bars != BA_EXPECTED_BARS:
        wrong = {code: (got, BA_EXPECTED_BARS.get(code))
                 for code, got in ba_bars.items()
                 if BA_EXPECTED_BARS.get(code) != got}
        failing("BA bars wrong (code: got, wanted) " + repr(wrong), failures)
    else:
        print("   OK  BA: default D, with C on the " +
              str(sum(1 for v in ba_bars.values() if v == "C")) +
              " major categories")

    # A Pass column must not be mistaken for a track. It is in
    # RESERVED_HEADINGS for exactly this reason; without that the file would
    # fail later, for a reason that reads as nonsense.
    if "Pass" in ba.tracks:
        failing("the Pass column was read as a track: " + repr(ba.tracks),
                failures)
    else:
        print("   OK  Pass is not a track (tracks: " +
              ", ".join(ba.tracks) + ")")
    return work, ba


def check_validation(work, failures):
    """A bar that is not a real grade is refused, with a usable message."""
    print()
    print("2. validation")

    for label, cell, value in (
            ("Requirements Pass cell", "F3", "F"),
            ("Requirements Pass cell", "F3", "banana"),
            ("Meta default", None, "F"),
    ):
        path = work / ("bad-" + str(abs(hash((label, value)))) + ".xlsx")
        workbook = check_support.workbook_from_tsv(
            check_support.BA_REQUIREMENTS_TSV)
        if cell:
            workbook["Requirements"][cell] = value
        else:
            meta = workbook["Meta"]
            for row in range(1, meta.max_row + 1):
                if str(meta.cell(row=row, column=1).value or "").lower() \
                        == "default passing grade":
                    meta.cell(row=row, column=2, value=value)
        workbook.save(str(path))

        try:
            requirements_io.load_requirements(path)
        except requirements_io.RequirementsError as error:
            # Quoted AS TYPED, not normalised — the person fixing the file
            # has to be able to find the cell by searching for what is in it.
            if repr(value) not in str(error):
                failing(label + " " + repr(value) + " was rejected, but the "
                        "message does not quote the offending value: " +
                        str(error), failures)
            else:
                print("   OK  " + label + " " + repr(value) + " rejected")
        else:
            failing(label + " " + repr(value) + " was accepted; F and stray "
                    "text rank 0 and would set the bar at zero", failures)


def check_tracker(work, ba, failures):
    """The generated tracker carries the right bars and the right lookup."""
    print()
    print("3. the generated tracker")

    tracker = work / "ba-tracker.xlsx"
    result = subprocess.run(
        [sys.executable, str(HERE / "transcript_to_tracker.py"),
         str(check_support.sample_transcript()),
         "-r", str(work / "ba-math-2020.xlsx"),
         "-o", str(tracker), "-t", "Honours"],
        capture_output=True, text=True,
    )
    if result.returncode != 0:
        sys.stderr.write(result.stdout + result.stderr)
        failing("could not generate a tracker against the BA file", failures)
        return None

    workbook = openpyxl.load_workbook(str(tracker))
    sheet = workbook["Tracker"]
    headings = {sheet.cell(row=5, column=c).value: c
                for c in range(1, sheet.max_column + 1)
                if sheet.cell(row=5, column=c).value}

    if "Pass" not in headings:
        failing("the Courses table has no Pass column", failures)
        return None
    if headings["Pass"] != headings["Gr."] + 1 \
            or headings["Ovr."] != headings["Pass"] + 1:
        failing("Pass should sit between Gr. and Ovr., got columns " +
                repr({k: headings[k] for k in ("Gr.", "Pass", "Ovr.")}),
                failures)
    else:
        print("   OK  Pass sits between Gr. and Ovr.")

    # The block's hidden bar column — Block.col_pass, SEVEN to the right of
    # the category column, which row 4 names. (Five until the Planned column
    # and its numeric helper landed between them.)
    block_col = None
    for column in range(1, sheet.max_column + 1):
        value = str(sheet.cell(row=4, column=column).value or "")
        if value.startswith("DEGREE REQUIREMENTS"):
            block_col = column
            break
    if block_col is None:
        failing("no requirements block found on row 4", failures)
        return None

    written = {}
    for row in range(6, 6 + len(ba.categories)):
        code = sheet.cell(row=row, column=block_col).value
        if code:
            written[code] = sheet.cell(row=row, column=block_col + 7).value
    if written != BA_EXPECTED_BARS:
        failing("the block's bar column holds " + repr(written), failures)
    else:
        print("   OK  every category row carries its bar (" +
              str(len(written)) + " categories)")

    if not sheet.column_dimensions[
            openpyxl.utils.get_column_letter(block_col + 7)].hidden:
        failing("the block's bar column is not hidden", failures)

    # The Pass cells must be a lookup, not a typed-in letter: the bar has to
    # follow the advisor when they recategorise a course.
    formula = sheet.cell(row=6, column=headings["Pass"]).value
    cat_letter = openpyxl.utils.get_column_letter(headings["Cat. 2020"])
    if not str(formula).startswith("=IFERROR(INDEX("):
        failing("the Pass cell is not a lookup: " + repr(formula), failures)
    elif "MATCH($" + cat_letter + "6," not in str(formula):
        failing("the Pass lookup does not key on this row's Cat. cell: " +
                repr(formula), failures)
    elif '"' + ba.default_pass_grade + '")' not in str(formula):
        failing("the Pass lookup does not fall back to the file default " +
                repr(ba.default_pass_grade) + ": " + repr(formula), failures)
    else:
        print("   OK  Pass is looked up from Cat., defaulting to " +
              ba.default_pass_grade)

    # An uncategorised row is the case the fallback exists for, and there are
    # always plenty: about 40% of a transcript needs a human.
    data = tracker_reader.read_tracker(tracker)
    blank_cat = [r for r in data["rows"] if not r.get("cat_2020")]
    if not blank_cat:
        failing("no uncategorised rows — the default bar is untested here",
                failures)
    else:
        bar = expected.bar_for(blank_cat[0], ba, "2020")
        if bar != ba.default_pass_grade:
            failing("an uncategorised row resolved to " + bar, failures)
        else:
            print("   OK  " + str(len(blank_cat)) + " uncategorised rows fall "
                  "back to " + bar)
    return data


def check_semantics(failures):
    """The rule itself, through the model the manual comparison relies on."""
    print()
    print("4. counting semantics")
    for grade, override, bar, should in SEMANTICS:
        row = {"grade": grade, "ovr": override}
        got = expected.counts_toward_progress(row, bar)
        described = "{:2s} {:6s} bar {}".format(grade or "—", override, bar)
        if got != should:
            failing(described + " -> " + str(got) + ", wanted " + str(should),
                    failures)
        else:
            print("   OK  " + described + " -> " +
                  ("counts" if got else "excluded"))


def main():
    failures = []
    work, ba = check_schema(failures)
    check_validation(work, failures)
    check_tracker(work, ba, failures)
    check_semantics(failures)

    print()
    if failures:
        print("FAILED —", len(failures), "problem(s) above")
        return 1
    print("PASSED — bars come from the requirements file, a D counts outside "
          "the major and not inside it.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
