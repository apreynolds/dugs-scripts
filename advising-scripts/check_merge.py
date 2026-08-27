# /// script
# dependencies = ["openpyxl", "pdfplumber", "reportlab"]
# ///
# check_merge.py — end-to-end test that a re-run preserves advisor edits.
# Usage: uv run check_merge.py
#
# A regression check, not part of the shipped pipeline. Simulates the yearly workflow:
#   1. generate a tracker from the transcript
#   2. an advisor hand-edits Cat., Ovr., Notes and a gpa1 flag, and renames
#      the student
#   3. the transcript is re-run against that tracker
#   4. every edit must survive, with no duplicated rows

import sys

import openpyxl

import check_support
import tracker_reader


def run(*args):
    return check_support.run_script("transcript_to_tracker.py", *args)


def main():
    WORK, PDF, REQ = check_support.prepare("_merge_test")
    tracker = WORK / "student.xlsx"

    print("STEP 1 — generate")
    run(str(PDF), "-r", str(REQ), "-o", str(tracker), "-t", "Honours")

    print("STEP 2 — simulate advisor edits")
    workbook = openpyxl.load_workbook(str(tracker))
    sheet = workbook["Tracker"]

    # Correct the student's identity (the parser only guesses at it).
    sheet["B1"] = "Gretzky, Wayne"
    sheet["B2"] = "3213213"

    edits = {}
    for row in range(6, 60):
        course = sheet.cell(row=row, column=2).value
        if course == "PHYS 1051":
            sheet.cell(row=row, column=7, value="SciLec")     # Cat. 2020
            sheet.cell(row=row, column=8, value="counts as first-year science")
            sheet.cell(row=row, column=9, value=True)          # gpa1
            edits["PHYS 1051"] = ("SciLec", "",
                                  "counts as first-year science", True)
        if course == "CS 1073":
            sheet.cell(row=row, column=6, value="no")          # Ovr.
            sheet.cell(row=row, column=7, value="FreeElec")
            edits["CS 1073"] = ("FreeElec", "no", "", False)
        if course == "MATH 3063":
            sheet.cell(row=row, column=7, value="Breadth1")
            sheet.cell(row=row, column=10, value=True)         # gpa2
            edits["MATH 3063"] = ("Breadth1", "", "", False)
    workbook.save(str(tracker))
    print("   edited", len(edits), "rows plus the name/ID")

    before = tracker_reader.read_tracker(tracker)
    n_before = len(before["rows"])

    print("STEP 3 — re-run the same transcript against it")
    output = run(str(PDF), "-r", str(REQ), "-u", str(tracker))
    print("   " + "\n   ".join(
        line for line in output.splitlines() if line.strip()))

    print("STEP 4 — verify")
    after = tracker_reader.read_tracker(tracker)
    by_course = {r["course"]: r for r in after["rows"]}

    failures = []

    if len(after["rows"]) != n_before:
        failures.append("row count changed: " + str(n_before) + " -> " +
                        str(len(after["rows"])))

    # Duplicates are keyed on (term, course), matching the merge key. The same
    # course in two different terms is a RETAKE, not a duplicate — the sample
    # transcript contains one deliberately.
    keys = [(r["term"], r["course"]) for r in after["rows"] if r["course"]]
    duplicates = sorted({k for k in keys if keys.count(k) > 1})
    if duplicates:
        failures.append("duplicated rows: " +
                        ", ".join(t + " " + c for t, c in duplicates))

    if after["student"]["name"] != "Gretzky, Wayne":
        failures.append("name not preserved: " +
                        repr(after["student"]["name"]))
    if after["student"]["id"] != "3213213":
        failures.append("id not preserved: " + repr(after["student"]["id"]))
    if after["track"] != "Honours":
        failures.append("track not preserved: " + repr(after["track"]))

    for course, (cat, ovr, notes, gpa1) in edits.items():
        row = by_course.get(course)
        if row is None:
            failures.append(course + " vanished")
            continue
        if row.get("cat_2020") != cat:
            failures.append(course + " Cat. = " + repr(row.get("cat_2020")) +
                            ", expected " + repr(cat))
        if row.get("ovr", "") != ovr:
            failures.append(course + " Ovr. = " + repr(row.get("ovr")) +
                            ", expected " + repr(ovr))
        if notes and row.get("notes", "") != notes:
            failures.append(course + " Notes = " + repr(row.get("notes")))
        if gpa1 and not row.get("gpa1"):
            failures.append(course + " gpa1 flag lost")

    if by_course.get("MATH 3063", {}).get("gpa2") is not True:
        failures.append("MATH 3063 gpa2 flag lost")

    if failures:
        print()
        print("FAILED:")
        for failure in failures:
            print("  -", failure)
        return 1
    print("   OK — all advisor edits survived, " + str(len(after["rows"])) +
          " rows, no duplicates.")

    # --- the harder cases ---------------------------------------------------

    print()
    print("STEP 5 — delete a whole term, and hand-add a planned course")
    workbook = openpyxl.load_workbook(str(tracker))
    sheet = workbook["Tracker"]

    # Blank out the 2026-1W rows, as though this transcript predated them.
    # NOTE: openpyxl's sheet.cell(row, column, value=None) is a no-op — the
    # value argument is only applied when it is not None — so the cell has to
    # be cleared by assigning to .value directly.
    removed = 0
    for row in range(6, 70):
        if sheet.cell(row=row, column=1).value == "2026-1W":
            for column in range(1, 11):
                sheet.cell(row=row, column=column).value = None
            removed += 1

    # A course the advisor planned by hand that no transcript will mention.
    first_blank = None
    for row in range(6, 70):
        if not sheet.cell(row=row, column=1).value:
            first_blank = row
            break
    sheet.cell(row=first_blank, column=1, value="2027-1W")
    sheet.cell(row=first_blank, column=2, value="MATH 4243")
    sheet.cell(row=first_blank, column=3, value="TOPICS IN ALGEBRA")
    sheet.cell(row=first_blank, column=4, value=3)
    sheet.cell(row=first_blank, column=7, value="MathElec")
    sheet.cell(row=first_blank, column=8, value="planned - not yet registered")
    workbook.save(str(tracker))
    print("   blanked " + str(removed) + " rows from 2026-1W, "
          "added a planned MATH 4243")

    print("STEP 6 — re-run")
    output = run(str(PDF), "-r", str(REQ), "-u", str(tracker))
    print("   " + "\n   ".join(
        line for line in output.splitlines() if line.strip()))

    print("STEP 7 — verify the harder cases")
    final = tracker_reader.read_tracker(tracker)
    final_by_course = {r["course"]: r for r in final["rows"]}
    late = []

    if len(final["rows"]) != n_before + 1:
        late.append("expected " + str(n_before + 1) + " rows (the transcript "
                    "plus the planned course), got " +
                    str(len(final["rows"])))

    restored = [r for r in final["rows"] if r["term"] == "2026-1W"]
    if len(restored) != removed:
        late.append("2026-1W not fully restored: " + str(len(restored)) +
                    " of " + str(removed))

    planned = final_by_course.get("MATH 4243")
    if planned is None:
        late.append("the hand-added MATH 4243 was deleted")
    else:
        if planned.get("cat_2020") != "MathElec":
            late.append("planned course lost its category")
        if planned.get("notes") != "planned - not yet registered":
            late.append("planned course lost its note")

    # Earlier edits must still be intact after this second merge.
    for course, (cat, _ovr, _notes, _gpa1) in edits.items():
        row = final_by_course.get(course)
        if row is None or row.get("cat_2020") != cat:
            late.append(course + " lost its category on the second merge")

    print()
    if late:
        print("FAILED:")
        for failure in late:
            print("  -", failure)
        return 1
    print("PASSED — deleted term re-added, hand-written planned course kept "
          "and reported, earlier edits still intact.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
