# programs.py — how each course figures in our degree programs, and reading class times.
# Usage: imported by timetable_grid.py; not run directly.
#
# Taken from check_timetable.py, the clash checker written for the 2027-28
# timetable committee and since retired: these are the parts the grids use.
# It uses build_courses.py and build_advising.py, so it works only when
# imported by a script beside them.
#
# HOW A COURSE'S PLACE IN A PROGRAM IS READ
# Each requirements workbook lists categories, each with the courses that
# satisfy it and the credit hours it needs on each track (Major, Honours).
# Courses not offered in the last three years (old numbers) are set aside.
#   required   the category needs as many hours as its courses add up to,
#              so a student takes them all (2ndMath, StatCore)
#   option     the category needs fewer hours than its courses offer, so a
#              student picks some (MathOpt, "MATH 3003 or 3103")
#   elective   the course is only caught by a pattern like MATH 3***
#
# A workbook's optional Year column says when each category's courses are
# suggested in the program ("2nd", "3rd/4th"); requirements_io.cell_year
# reads it, and suggested_years() looks up a course's.

import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

import build_advising as ba
import build_courses as bc
import requirements_io  # from lib/, beside this file

OURS = ("MATH", "STAT", "DATA")

# How many past falls (or winters) to search for another unit's time.
LOOKBACK = 3
# How many years back a course must have run to count as still offered.
LIVE_YEARS = 3

# Courses a requirements workbook lists beside the course they stand in for,
# {substitute: usual course}. Taken from the workbooks' own Description and
# Note cells ("MATH 1053/1063 accepted in place of 1003/1013", "PHYS 1082 may
# replace PHYS 1052", "Or ECON 1014/1024", "Or MATH 3503"). Without this a
# category like Calc looks like "any two of four", and its usual courses read
# as options rather than requirements.
SUBSTITUTES = {
    "MATH1053": "MATH1003", "MATH1063": "MATH1013",
    "PHYS1082": "PHYS1052",
    "ECON1014": "ECON1013", "ECON1024": "ECON1023",
    "MATH3503": "MATH3043",
}

ROLE_RANK = {"required": 3, "option": 2, "elective": 1}

DAY = re.compile(r"Th|Sa|Su|M|T|W|F|S")
MEETING = re.compile(r"([A-Za-z/]+)\s+(\d{1,2}):(\d{2})\s*-\s*(\d{1,2}):(\d{2})")
# 'MATH 3***'. It ends with a lookahead, not \b: there is no word boundary
# between '*' and ';' or the end of the text, so \b never matched a pattern.
WILDCARD = re.compile(r"\b([A-Z]{2,4})\s?(\d[\d*]{3})(?![\d*])")


# ----------------------------------------------------------------------------
# Times
# A meeting is (day, start, end) in minutes after midnight. An option is one
# section's meetings, e.g. MWF 10:30-11:20 is three meetings.
# ----------------------------------------------------------------------------
def clock(hours, mins):
    """'1', '30' -> 13:30. Times carry no AM/PM, and no class starts before 8."""
    hours = int(hours)
    if hours < 8:
        hours += 12
    return hours * 60 + int(mins)


def parse_meetings(text):
    """'MWF 10:30-11:20' or 'M 2:30-3:50, W 2:30-3:50' -> a list of meetings."""
    meetings = []
    for days, h1, m1, h2, m2 in MEETING.findall(text):
        for day in DAY.findall(days.replace("/", "")):
            meetings.append((day, clock(h1, m1), clock(h2, m2)))
    return meetings


def parse_options(text):
    """A time as written -> options. 'and' or '|' separates sections at different times."""
    text = re.sub(r"\(.*?\)", "", str(text or ""))  # '(set by Economics)'
    options = [parse_meetings(part) for part in re.split(r"\band\b|\|", text)]
    return [option for option in options if option]


# ----------------------------------------------------------------------------
# Reading the inputs
# ----------------------------------------------------------------------------
def read_proposal(path):
    """A proposed timetable's term sheets -> [(sheet title, 'FA' or 'WI', {course: row})]."""
    book = load_workbook(path, data_only=True)
    terms = []
    for sheet in book.worksheets:
        match = re.fullmatch(r"(Fall|Winter) (\d{4})", sheet.title)
        if not match:
            continue
        rows = sheet.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
        courses = {}
        for row in rows:
            record = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, row) if h}
            key = bc.course_key(record.get("Course", ""))
            if ba.any_course(key):
                courses[key] = record
        terms.append((sheet.title, "FA" if match[1] == "Fall" else "WI", courses))
    if not terms:
        sys.exit(f"{path.name} has no 'Fall YYYY' or 'Winter YYYY' sheet.")
    return terms


def read_aliases(path):
    """{other number: the course it belongs to}, from the advising timetable's Other numbers.

    ECON 4673 -> MATH 3373 (cross-listed), PHYS 1061 -> PHYS 1051 (renumbered).
    That workbook is where these are recorded for every subject, so it is read
    rather than duplicated here.
    """
    book = load_workbook(path, data_only=True)
    aliases = {}
    for key, record in ba.sheet_records(book["Timetable"]).items():
        for other in ba.codes_in(record.get("Other numbers", "")):
            aliases.setdefault(other, key)
    return aliases


def read_programs(folder, aliases, live, hours_of):
    """Read every requirements workbook into the role each course plays.

    Returns a list of programs, each {"name", "title", "tracks": {track: {"exact":
    {course: [(role, category, need)]}, "wild": [(pattern, category)]}},
    "years": {category: (year, ...)}}.
    need is roughly how many of the category's courses a student takes;
    years are the Year column's, e.g. (3, 4) for "3rd/4th", and only the
    categories with one are in it.
    """
    programs = []
    for path in sorted(Path(folder).glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue  # Excel's lock file, not a workbook
        book = load_workbook(path, data_only=True)
        if "Requirements" not in book.sheetnames or "Meta" not in book.sheetnames:
            continue
        # The tracks are the "Total CH: <track>" rows of Meta; the Requirements
        # sheet has a column of hours for each.
        tracks = []
        title = path.stem
        for field, value, *_ in book["Meta"].iter_rows(values_only=True):
            if str(field or "").startswith("Total CH: "):
                tracks.append(str(field)[len("Total CH: "):].strip())
            elif field == "Program" and value:
                title = str(value).strip()  # 'BSc Mathematics', for headings
        rows = book["Requirements"].iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
        program = {"name": re.sub(r"-\d{4}$", "", path.stem), "title": title,
                   "tracks": {t: {"exact": defaultdict(list), "wild": []} for t in tracks},
                   "years": {}}
        # The Year heading, in whatever case it was typed.
        year_heading = next((h for h in header if h.lower() == "year"), None)
        for number, row in enumerate(rows, start=2):
            record = dict(zip(header, row))
            category = str(record.get("Category") or "").strip()
            if year_heading:
                years = requirements_io.cell_year(record.get(year_heading), f"Requirements row {number} Year", path)
                if years:
                    program["years"][category] = years
            text = str(record.get("Courses") or "")
            exact = []
            for subject, number in ba.CODE.findall(text):
                key = aliases.get(subject + number, subject + number)
                if key not in exact:
                    exact.append(key)
            patterns = [re.compile(subject + number.replace("*", r"\d"))
                        for subject, number in WILDCARD.findall(text) if "*" in number]
            current = [c for c in exact if c in live]
            # A substitute listed beside its usual course is left out of the
            # sum, so "MATH 1003; MATH 1053; MATH 1013; MATH 1063" for 6 ch
            # reads as MATH 1003 and 1013 required, not "two of four".
            subs = [c for c in current if SUBSTITUTES.get(c) in current]
            main_courses = [c for c in current if c not in subs]
            for track in tracks:
                try:
                    hours = float(record.get(track) or 0)
                except ValueError:
                    hours = 0.0  # '-' marks the slack category
                if hours <= 0:
                    continue
                entry = program["tracks"][track]
                if main_courses:
                    total = sum(hours_of.get(c, 3.0) for c in main_courses)
                    role = "required" if total <= hours + 0.01 else "option"
                    need = max(1, round(hours / (total / len(main_courses))))
                    for course in main_courses:
                        entry["exact"][course].append((role, category, need))
                    for course in subs:
                        entry["exact"][course].append(("option", category, 1))
                for pattern in patterns:
                    entry["wild"].append((pattern, category))
        programs.append(program)
    return programs


def load_programs(requirements, advising, timetables, extra_live=()):
    """Read everything a program check needs.

    Returns (programs, offerings, covered terms in order, timetable files,
    the set of courses still offered).
    extra_live: courses to treat as still offered whatever the exports say,
    such as everything on the proposed timetable.
    """
    aliases = read_aliases(Path(advising).expanduser())
    counted = defaultdict(list)
    for other, key in aliases.items():
        counted[key].append(other)
    timetable, covered, tt_files = bc.load_timetables(timetables, subjects=None)
    offerings = bc.build_offerings(timetable, covered, {}, [], counted, {})
    terms = sorted(covered, key=bc.term_key)
    # Still offered: any section, in any mode, in the last LIVE_YEARS years.
    recent = set(terms[-3 * LIVE_YEARS:])
    live = {c for c, per in offerings.items() if recent & set(per)} | set(extra_live)
    hours_of = {}
    for (course, term), rows in sorted(timetable.items(), key=lambda kv: bc.term_key(kv[0][1])):
        try:
            hours_of[aliases.get(course, course)] = float(rows[0]["Credit Hours"])
        except ValueError:
            pass
    programs = read_programs(requirements, aliases, live, hours_of)
    return programs, offerings, terms, tt_files, live


def suggested_years(program, course):
    """The years a program suggests a course for, from its categories' Year cells: e.g. (2,).

    The categories that name the course outright, on any track; if none does,
    those whose pattern (MATH 3***) catches it, as in the tracker. A course in
    two categories is suggested for the years of both. () when no category has
    a year for it (not in one, or "-"): the caller falls back on its level.
    """
    tracks = program["tracks"].values()
    categories = {category for entry in tracks for _, category, _ in entry["exact"].get(course, [])}
    if not categories:
        categories = {category for entry in tracks for pattern, category in entry["wild"]
                      if pattern.fullmatch(course)}
    return tuple(sorted({year for c in categories for year in program["years"].get(c, ())}))


def role_in(track_entry, course):
    """The strongest role a course has on one track: (role, category, need), or None."""
    named = track_entry["exact"].get(course)
    if named:
        return max(named, key=lambda r: ROLE_RANK[r[0]])
    # An exact entry beats a pattern, as in the tracker.
    for pattern, category in track_entry["wild"]:
        if pattern.fullmatch(course):
            return ("elective", category, 0)
    return None
