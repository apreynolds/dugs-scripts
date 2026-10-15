#!/usr/bin/env -S uv run --script
# /// script
# dependencies = ["openpyxl"]
# ///
# check_timetable.py — check a proposed timetable against the degree programs.
# Usage:
#   ./check_timetable.py                    check, and write the findings into the workbook
#   ./check_timetable.py -n                 check and print only; write nothing
#   ./check_timetable.py --try "MATH 3213=TTh 1:00-2:20" [--try ...]
#   ./check_timetable.py -h | --help
#
#   -w, --workbook      the proposed timetable (default: proposed-timetable-2027-28.xlsx
#                       in timetable-committee). Its "Fall YYYY" and "Winter YYYY"
#                       sheets are read; only the three "Check - ..." sheets are
#                       written, replacing the ones from the last run. A dated
#                       copy of the workbook goes into a zzzzz folder first.
#                       Refuses to run while the workbook is open in Excel
#   -r, --requirements  folder of degree-requirements workbooks
#   -a, --advising      the advising timetable, read only for its Other numbers
#                       (cross-listings and renumberings)
#   -t, --timetable     an Informer timetable export, or a directory of them;
#                       repeat the flag for several
#       --try C=TIME    also test course C at TIME, as if it were one more
#                       alternative; repeat for several. Not written anywhere
#                       but the Alternatives sheet and the screen
#   -n, --dry-run       print the findings and write nothing
#
# WHAT THIS IS FOR
# The timetable committee rolls each course over into last year's slot, then
# looks for clashes. A clash matters when a student needs both courses in the
# same term, and that depends on the programs. This script answers, for every
# pair of courses that would meet at the same time: which programs name both,
# and how firmly (both required, or one of them only an option).
# It is meant to be rerun whenever a time changes in the workbook.
#
# WHERE THE TIMES COME FROM
#   our courses        the workbook: "Suggested time", "Alternative time" and
#                      "Tutorial". A tutorial is given only as a start time,
#                      so it is taken to last 50 minutes
#   other units'       the timetable exports: the most recent fall (or winter)
#   courses            in which the course ran in person, looking back at most
#                      three of them. Their 2027-28 times are not known yet
#
# WHAT COUNTS AS A CLASH
# Two courses clash only if a student cannot dodge it: every section of one
# overlaps every section of the other. A course with a second lecture or a
# choice of labs at a free time does not clash. Lectures and labs/tutorials
# are both checked, so a lecture landing on every lab section counts.
#
# HOW A COURSE'S PLACE IN A PROGRAM IS READ
# Each requirements workbook lists categories, each with the courses that
# satisfy it and the credit hours it needs on each track (Major, Honours).
# Courses not offered in the last three years (old numbers) are set aside.
#   required   the category needs as many hours as its courses add up to,
#              so a student takes them all (2ndMath, StatCore)
#   option     the category needs fewer hours than its courses offer, so a
#              student picks some (MathOpt, "MATH 3003 or 3103")
#   elective   the course is only caught by a pattern like MATH 3***
# Two options in the same category that needs only one course are
# alternatives to each other, so a clash between them does not matter.
# Each pair is graded by the worst case over all programs and tracks:
#   Both required, Required + option, Both options, Required + elective
#   (an elective caught by a pattern like MATH 3*** against a required
#   course), Elective (an elective against an option or another elective), and
#   Ours only (two of our courses that no program names together; this
#   is what the "no two of our courses overlap" guideline is about).
#
# This script borrows its timetable reading from build_courses.py and its
# requirements reading from build_advising.py, which must sit in the same folder.

import argparse
import datetime
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

import build_advising as ba
import build_courses as bc

# ----------------------------------------------------------------------------
# Defaults
# ----------------------------------------------------------------------------
DEFAULT_WORKBOOK = bc.DUGS / "timetable-committee/proposed-timetable-2027-28.xlsx"

OURS = ("MATH", "STAT", "DATA")

# How many past falls (or winters) to search for another unit's time.
LOOKBACK = 3
# How many years back a course must have run to count as still offered.
LIVE_YEARS = 3
# A tutorial in the workbook has only a start time.
TUTORIAL_MINUTES = 50

# The sheets this script owns in the workbook. Everything else is Patrick's.
SHEET_CLASHES = "Check - clashes"
SHEET_ALTERNATIVES = "Check - alternatives"
SHEET_ABOUT = "Check - about"

# Courses a requirements workbook lists beside the course they stand in for,
# {substitute: usual course}. Taken from the workbooks' own Description and
# Note cells ("MATH 1053/1063 accepted in place of 1003/1013", "PHYS 1082 may
# replace PHYS 1052", "Or ECON 1014/1024", "Or MATH 3503"). Without this a
# category like Calc looks like "any two of four", and its usual courses read
# as options rather than requirements.
SUBSTITUTES = {
    "MATH1053": "MATH1003", "MATH1063": "MATH1013",
    "PHYS1082": "PHYS1052",
    "ECON1014": "ECON1013", "ECON1024": "ECON1023",
    "MATH3503": "MATH3043",
}

# Worst first.
SEVERITY = ["Both required", "Required + option", "Both options", "Required + elective", "Elective", "Ours only"]
ROLE_RANK = {"required": 3, "option": 2, "elective": 1}

DAY = re.compile(r"Th|Sa|Su|M|T|W|F|S")
DAY_ORDER = ["M", "T", "W", "Th", "F", "Sa", "Su", "S"]
MEETING = re.compile(r"([A-Za-z/]+)\s+(\d{1,2}):(\d{2})\s*-\s*(\d{1,2}):(\d{2})")
# 'MATH 3***'. It ends with a lookahead, not \b: there is no word boundary
# between '*' and ';' or the end of the text, so \b never matched a pattern.
WILDCARD = re.compile(r"\b([A-Z]{2,4})\s?(\d[\d*]{3})(?![\d*])")


# ----------------------------------------------------------------------------
# Times
# A meeting is (day, start, end) in minutes after midnight. An option is one
# section's meetings, e.g. MWF 10:30-11:20 is three meetings. A component is
# the list of options a student chooses one from: the lecture sections, or
# the lab/tutorial sections.
# ----------------------------------------------------------------------------
def clock(hours, mins):
    """'1', '30' -> 13:30. Times carry no AM/PM, and no class starts before 8."""
    hours = int(hours)
    if hours < 8:
        hours += 12
    return hours * 60 + int(mins)


def parse_meetings(text):
    """'MWF 10:30-11:20' or 'M 2:30-3:50, W 2:30-3:50' -> a list of meetings."""
    meetings = []
    for days, h1, m1, h2, m2 in MEETING.findall(text):
        for day in DAY.findall(days.replace("/", "")):
            meetings.append((day, clock(h1, m1), clock(h2, m2)))
    return meetings


def parse_options(text):
    """A workbook time -> options. 'and' or '|' separates sections at different times."""
    text = re.sub(r"\(.*?\)", "", str(text or ""))  # '(set by Economics)'
    options = [parse_meetings(part) for part in re.split(r"\band\b|\|", text)]
    return [option for option in options if option]


def parse_tutorials(text):
    """The workbook's Tutorial cell, 'T, Th 8:30' -> one option per day, 50 minutes each."""
    text = re.sub(r"\(.*?\)", "", str(text or "")).strip()
    match = re.fullmatch(r"([A-Za-z,/ ]+?)\s+(\d{1,2}):(\d{2})", text)
    if not match:
        return parse_options(text)  # written out in full, or blank
    start = clock(match[2], match[3])
    return [[(day, start, start + TUTORIAL_MINUTES)] for day in DAY.findall(match[1].replace(" ", ""))]


def overlap(a, b):
    """Minutes two options share, by day: {'M': 50, 'W': 50, 'F': 50}."""
    shared = {}
    for day, s1, e1 in a:
        for day2, s2, e2 in b:
            if day == day2 and s1 < e2 and s2 < e1:
                shared[day] = max(shared.get(day, 0), min(e1, e2) - max(s1, s2))
    return shared


def forced_overlap(component_a, component_b):
    """The overlap if no choice of sections avoids it, else None."""
    if not component_a or not component_b:
        return None
    first = None
    for a in component_a:
        for b in component_b:
            shared = overlap(a, b)
            if not shared:
                return None  # this pair of sections is free, so the clash can be dodged
            first = first or shared
    return first


def clash(timing_a, timing_b):
    """(which meetings, overlap) for two courses' timings, or None.

    A timing is {"Lecture": options, "Lab/tut": options}. Lectures against
    lectures are checked first, since that is the clash people mean.
    """
    for part_a in ("Lecture", "Lab/tut"):
        for part_b in ("Lecture", "Lab/tut"):
            shared = forced_overlap(timing_a.get(part_a), timing_b.get(part_b))
            if shared:
                kind = part_a if part_a == part_b else f"{part_a} / {part_b}"
                return kind.lower(), shared
    return None


def overlap_text(shared):
    """{'M': 50, 'W': 50, 'F': 50} -> 'MWF 50 min'; mixed lengths are listed day by day."""
    days = sorted(shared, key=DAY_ORDER.index)
    if len(set(shared.values())) == 1:
        return f"{''.join(days)} {shared[days[0]]} min"
    return ", ".join(f"{day} {shared[day]} min" for day in days)


# ----------------------------------------------------------------------------
# Reading the inputs
# ----------------------------------------------------------------------------
def read_proposal(path):
    """The workbook's term sheets -> [(sheet title, 'FA' or 'WI', {course: row})]."""
    book = load_workbook(path, data_only=True)
    terms = []
    for sheet in book.worksheets:
        match = re.fullmatch(r"(Fall|Winter) (\d{4})", sheet.title)
        if not match:
            continue
        rows = sheet.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
        courses = {}
        for row in rows:
            record = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, row) if h}
            key = bc.course_key(record.get("Course", ""))
            if ba.any_course(key):
                courses[key] = record
        terms.append((sheet.title, "FA" if match[1] == "Fall" else "WI", courses))
    if not terms:
        sys.exit(f"{path.name} has no 'Fall YYYY' or 'Winter YYYY' sheet.")
    return terms


def read_aliases(path):
    """{other number: the course it belongs to}, from the advising timetable's Other numbers.

    ECON 4673 -> MATH 3373 (cross-listed), PHYS 1061 -> PHYS 1051 (renumbered).
    That workbook is where these are recorded for every subject, so it is read
    rather than duplicated here.
    """
    book = load_workbook(path, data_only=True)
    aliases = {}
    for key, record in ba.sheet_records(book["Timetable"]).items():
        for other in ba.codes_in(record.get("Other numbers", "")):
            aliases.setdefault(other, key)
    return aliases


def read_programs(folder, aliases, live, hours_of):
    """Read every requirements workbook into the role each course plays.

    Returns a list of programs, each {"name", "title", "tracks": {track: {"exact":
    {course: [(role, category, need)]}, "wild": [(pattern, category)]}}}.
    need is roughly how many of the category's courses a student takes.
    """
    programs = []
    for path in sorted(Path(folder).glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue  # Excel's lock file, not a workbook
        book = load_workbook(path, data_only=True)
        if "Requirements" not in book.sheetnames or "Meta" not in book.sheetnames:
            continue
        # The tracks are the "Total CH: <track>" rows of Meta; the Requirements
        # sheet has a column of hours for each.
        tracks = []
        title = path.stem
        for field, value, *_ in book["Meta"].iter_rows(values_only=True):
            if str(field or "").startswith("Total CH: "):
                tracks.append(str(field)[len("Total CH: "):].strip())
            elif field == "Program" and value:
                title = str(value).strip()  # 'BSc Mathematics', for headings
        rows = book["Requirements"].iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
        program = {"name": re.sub(r"-\d{4}$", "", path.stem), "title": title,
                   "tracks": {t: {"exact": defaultdict(list), "wild": []} for t in tracks}}
        for row in rows:
            record = dict(zip(header, row))
            category = str(record.get("Category") or "").strip()
            text = str(record.get("Courses") or "")
            exact = []
            for subject, number in ba.CODE.findall(text):
                key = aliases.get(subject + number, subject + number)
                if key not in exact:
                    exact.append(key)
            patterns = [re.compile(subject + number.replace("*", r"\d"))
                        for subject, number in WILDCARD.findall(text) if "*" in number]
            current = [c for c in exact if c in live]
            # A substitute listed beside its usual course is left out of the
            # sum, so "MATH 1003; MATH 1053; MATH 1013; MATH 1063" for 6 ch
            # reads as MATH 1003 and 1013 required, not "two of four".
            subs = [c for c in current if SUBSTITUTES.get(c) in current]
            main_courses = [c for c in current if c not in subs]
            for track in tracks:
                try:
                    hours = float(record.get(track) or 0)
                except ValueError:
                    hours = 0.0  # '-' marks the slack category
                if hours <= 0:
                    continue
                entry = program["tracks"][track]
                if main_courses:
                    total = sum(hours_of.get(c, 3.0) for c in main_courses)
                    role = "required" if total <= hours + 0.01 else "option"
                    need = max(1, round(hours / (total / len(main_courses))))
                    for course in main_courses:
                        entry["exact"][course].append((role, category, need))
                    for course in subs:
                        entry["exact"][course].append(("option", category, 1))
                for pattern in patterns:
                    entry["wild"].append((pattern, category))
        programs.append(program)
    return programs


def load_programs(requirements, advising, timetables, extra_live=()):
    """Read everything a program check needs.

    Returns (programs, offerings, covered terms in order, timetable files,
    the set of courses still offered).
    extra_live: courses to treat as still offered whatever the exports say,
    such as everything on the proposed timetable.
    """
    aliases = read_aliases(Path(advising).expanduser())
    counted = defaultdict(list)
    for other, key in aliases.items():
        counted[key].append(other)
    timetable, covered, tt_files = bc.load_timetables(timetables, subjects=None)
    offerings = bc.build_offerings(timetable, covered, {}, [], counted, {})
    terms = sorted(covered, key=bc.term_key)
    # Still offered: any section, in any mode, in the last LIVE_YEARS years.
    recent = set(terms[-3 * LIVE_YEARS:])
    live = {c for c, per in offerings.items() if recent & set(per)} | set(extra_live)
    hours_of = {}
    for (course, term), rows in sorted(timetable.items(), key=lambda kv: bc.term_key(kv[0][1])):
        try:
            hours_of[aliases.get(course, course)] = float(rows[0]["Credit Hours"])
        except ValueError:
            pass
    programs = read_programs(requirements, aliases, live, hours_of)
    return programs, offerings, terms, tt_files, live


def role_in(track_entry, course):
    """The strongest role a course has on one track: (role, category, need), or None."""
    named = track_entry["exact"].get(course)
    if named:
        return max(named, key=lambda r: ROLE_RANK[r[0]])
    # An exact entry beats a pattern, as in the tracker.
    for pattern, category in track_entry["wild"]:
        if pattern.fullmatch(course):
            return ("elective", category, 0)
    return None


def pair_severity(programs, a, b):
    """How much a clash between a and b matters: (severity, text naming the programs)."""
    worst = len(SEVERITY) - 1  # "Ours only"
    parts = []
    for program in programs:
        found = {}  # severity index -> (tracks, description)
        for track, entry in program["tracks"].items():
            role_a, role_b = role_in(entry, a), role_in(entry, b)
            if not role_a or not role_b:
                continue
            roles = {role_a[0], role_b[0]}
            # Two options of one category that needs a single course ("MATH
            # 3003 or 3103") are alternatives: nobody needs both.
            if role_a[1] == role_b[1] and roles == {"option"} and role_a[2] < 2:
                continue
            if "elective" in roles:
                # A required course against an elective blocks that elective
                # for everyone taking the required one that year: worse than
                # two electives meeting.
                level = 3 if "required" in roles else 4
            elif roles == {"required"}:
                level = 0
            elif "required" in roles:
                level = 1
            else:
                level = 2
            text = f"{role_a[1]} ({role_a[0]}) × {role_b[1]} ({role_b[0]})"
            tracks, _ = found.get(level, ([], text))
            found[level] = (tracks + [track], text)
        if not found:
            continue
        level = min(found)
        tracks, text = found[level]
        label = "" if len(tracks) == len(program["tracks"]) else f" ({', '.join(tracks)})"
        parts.append((level, f"{program['name']}{label}: {text}"))
        worst = min(worst, level)
    if not parts and not (a[:-4] in OURS and b[:-4] in OURS):
        return None, ""  # another unit's course that no program puts with ours
    parts.sort()
    return SEVERITY[worst], "; ".join(text for _, text in parts)


def stage(key):
    """1, 2, or 3 for upper year (3000- and 4000-level courses are taken together).

    MATH 0863, a pre-university course, counts as first year.
    """
    return min(max(int(key[-4]), 1), 3)


def years_text(a, b):
    names = {1: "1st year", 2: "2nd year", 3: "upper"}
    if stage(a) == stage(b):
        return "same"
    return f"{names[stage(a)]} vs {names[stage(b)]}"


# ----------------------------------------------------------------------------
# The check
# ----------------------------------------------------------------------------
def our_timing(lecture_text, tutorial_text):
    timing = {"Lecture": parse_options(lecture_text)}
    tutorials = parse_tutorials(tutorial_text)
    if tutorials:
        timing["Lab/tut"] = tutorials
    return timing


def other_timing(per, terms):
    """Another unit's course: its timing in the latest of terms it ran in person, and its history."""
    history = []
    timing, used = None, None
    for term in reversed(terms):
        offering = per.get(term)
        if not (offering and offering["inperson"]):
            history.append(f"{bc.term_label(term)}: not offered")
            continue
        slots = " | ".join(offering["lecture_slots"])
        labs = offering["tutorial_slots"]
        history.append(f"{bc.term_label(term)}: {slots or 'labs only'}")
        if timing is None and (offering["lecture_slots"] or labs):
            timing = {}
            if offering["lecture_slots"]:
                timing["Lecture"] = [parse_meetings(s) for s in offering["lecture_slots"]]
            if labs:
                timing["Lab/tut"] = [parse_meetings(s) for s in labs]
            used = term
    return timing, used, "; ".join(history)


def find_clashes(term_courses, timings, programs, severity_cache):
    """Every clash in one term, given a timing for each course. One record per pair."""
    found = []
    keys = sorted(timings, key=lambda k: (k[:-4] not in OURS, k[:-4], k[-4:]))
    for i, a in enumerate(keys):
        if a not in term_courses:
            continue  # pairs always start from one of our courses
        for b in keys[i + 1:]:
            pair = (a, b)
            if pair not in severity_cache:
                severity_cache[pair] = pair_severity(programs, a, b)
            severity, programs_text = severity_cache[pair]
            if severity is None:
                continue
            result = clash(timings[a], timings[b])
            if result:
                kind, shared = result
                found.append({"a": a, "b": b, "kind": kind, "overlap": overlap_text(shared),
                              "severity": severity, "programs": programs_text})
    return found


def main():
    parser = argparse.ArgumentParser(
        description="Check a proposed timetable against the degree-requirements workbooks."
    )
    parser.add_argument("-w", "--workbook", default=DEFAULT_WORKBOOK,
                        help="the proposed timetable (default: the 2027-28 one)")
    parser.add_argument("-r", "--requirements", default=ba.DEFAULT_REQUIREMENTS,
                        help="folder of degree-requirements workbooks")
    parser.add_argument("-a", "--advising", default=ba.DEFAULT_OUTPUT,
                        help="the advising timetable, for other numbers")
    parser.add_argument("-t", "--timetable", action="append",
                        help="timetable export or directory of them (repeatable)")
    parser.add_argument("--try", dest="tries", action="append", default=[], metavar="COURSE=TIME",
                        help='also test a course at another time, e.g. "MATH 3213=TTh 1:00-2:20"')
    parser.add_argument("-n", "--dry-run", action="store_true",
                        help="print the findings and write nothing")
    args = parser.parse_args()

    workbook_path = Path(args.workbook).expanduser()
    if not args.dry_run:
        bc.refuse_if_open(workbook_path)
    today = datetime.date.today().isoformat()

    tries = defaultdict(list)
    for item in args.tries:
        course, _, time = item.partition("=")
        if not parse_options(time):
            sys.exit(f"--try {item!r}: write it as COURSE=TIME, e.g. \"MATH 3213=TTh 1:00-2:20\".")
        tries[bc.course_key(course)].append(time.strip())

    # ---- The data -----------------------------------------------------------
    proposal = read_proposal(workbook_path)
    extra = set()
    for _, _, courses in proposal:
        extra |= set(courses)
    programs, offerings, terms, tt_files, live = load_programs(
        args.requirements, args.advising, args.timetable or bc.DEFAULT_TIMETABLES, extra)
    named = set()
    for program in programs:
        for entry in program["tracks"].values():
            named |= set(entry["exact"])

    # ---- Check each term ----------------------------------------------------
    clash_rows, alternative_rows, unreadable = [], [], []
    severity_cache = {}
    checked = set()  # other units' courses that had a recent time in some term
    for title, season, courses in proposal:
        season_terms = [t for t in terms if t.endswith(season)][-LOOKBACK:]
        # Another unit's course takes the time it last had in this season.
        others, histories = {}, {}
        for course in sorted(named):
            if course[:-4] in OURS:
                continue
            timing, used, history = other_timing(offerings.get(course, {}), season_terms)
            if timing:
                others[course] = timing
                histories[course] = f"{history} (used {bc.term_label(used)})"
                checked.add(course)
        suggested, alternative, texts = {}, {}, {}
        for course, record in courses.items():
            timing = our_timing(record.get("Suggested time"), record.get("Tutorial"))
            if not timing["Lecture"]:
                unreadable.append(f"{title}: {bc.course_display(course)} {record.get('Suggested time')!r}")
                continue
            suggested[course] = timing
            texts[course] = record.get("Suggested time", "")
            alt = record.get("Alternative time", "")
            alternative[course] = our_timing(alt, record.get("Tutorial")) if parse_options(alt) else timing

        # Two scenarios: everything at its suggested time, and every
        # alternative taken at once. A row says which of them it occurs in.
        scenarios = {"Suggested": {**others, **suggested}, "All alternatives": {**others, **alternative}}

        def time_text(course, scenario):
            """The time a course has in a scenario, as written in the workbook."""
            if course not in courses:
                return other_time_text(histories[course])
            alt = courses[course].get("Alternative time", "")
            return alt if scenario == "All alternatives" and parse_options(alt) else texts[course]

        merged = {}
        for name, timings in scenarios.items():
            for found in find_clashes(courses, timings, programs, severity_cache):
                a, b = found["a"], found["b"]
                key = (a, b, time_text(a, name), time_text(b, name))
                merged.setdefault(key, {**found, "in": set()})["in"].add(name)
        for (a, b, time_a, time_b), found in merged.items():
            clash_rows.append([
                title, bc.course_display(a), time_a, bc.course_display(b), time_b,
                found["overlap"], found["kind"], found["severity"], years_text(a, b), found["programs"],
                "yes" if "Suggested" in found["in"] else "", "yes" if "All alternatives" in found["in"] else "",
                histories.get(b, ""),
            ])

        # Each course with an alternative (or a --try) at each of its times,
        # with everything else at its suggested time.
        for course, record in courses.items():
            candidates = [("Suggested", record.get("Suggested time", ""))]
            if parse_options(record.get("Alternative time")):
                candidates.append(("Alternative", record.get("Alternative time")))
            candidates += [("--try", t) for t in tries.get(course, [])]
            if len(candidates) == 1 or course not in suggested:
                continue
            for which, time in candidates:
                timings = {**others, **suggested, course: our_timing(time, record.get("Tutorial"))}
                found = [f for f in find_clashes(courses, timings, programs, severity_cache)
                         if course in (f["a"], f["b"])]
                found.sort(key=lambda f: (SEVERITY.index(f["severity"]), years_text(f["a"], f["b"]) != "same"))
                listed = []
                for f in found:
                    partner = f["b"] if f["a"] == course else f["a"]
                    note = "" if years_text(f["a"], f["b"]) == "same" else ", different years"
                    listed.append(f"{bc.course_display(partner)} ({f['severity']}{note}; {f['overlap']})")
                worst = found[0]["severity"] if found else "none"
                alternative_rows.append([title, bc.course_display(course), which, time, worst,
                                         len(found), "\n".join(listed)])
    for course in tries:
        if not any(course in courses for _, _, courses in proposal):
            print(f"  --try {bc.course_display(course)}: not on any term sheet, so not tested")

    clash_rows.sort(key=lambda r: (r[0][:4] != "Fall", SEVERITY.index(r[7]), r[8] != "same", r[1], r[3]))
    # A course that runs only in the winter is simply not in the fall check.
    # Only one with no recent in-person time in either term is worth reporting.
    no_time = [bc.course_display(c) for c in sorted(named)
               if c[:-4] not in OURS and c in live and c not in checked]

    # ---- Report -------------------------------------------------------------
    report(clash_rows, alternative_rows, unreadable, no_time)
    if args.dry_run:
        print("Dry run: nothing written.")
        return

    book = load_workbook(workbook_path)  # formulas kept, unlike data_only
    for name in (SHEET_CLASHES, SHEET_ALTERNATIVES, SHEET_ABOUT):
        if name in book.sheetnames:
            del book[name]
    columns = [
        ("Term", 12, "id"), ("Course", 11, "id"), ("Time", 22, "other"),
        ("Clashes with", 12, "id"), ("Its time", 24, "other"), ("Overlap", 16, "other"),
        ("Meetings", 16, "other"), ("Worst case", 17, "other"), ("Years", 18, "other"),
        ("Programs", 60, "other"), ("Suggested", 10, "other"), ("All alternatives", 11, "other"),
        ("Its recent times", 60, "other"),
    ]
    bc.write_table(book.create_sheet(SHEET_CLASHES), columns, clash_rows, freeze="C2")
    columns = [
        ("Term", 12, "id"), ("Course", 11, "id"), ("Which", 12, "other"), ("Time", 22, "other"),
        ("Worst case", 17, "other"), ("Clashes", 9, "other"), ("Clashes with", 80, "other"),
    ]
    bc.write_table(book.create_sheet(SHEET_ALTERNATIVES), columns, alternative_rows, freeze="C2")
    about_sheet(book.create_sheet(SHEET_ABOUT), today, args, tt_files, terms, programs)

    backup = bc.save_safely(book, workbook_path)
    print(f"Wrote the three 'Check - ...' sheets into {workbook_path}")
    if backup:
        print(f"  The copy it replaced was kept as {backup}")


def other_time_text(history):
    """Another unit's course: the time used, read back from its history ('... (used 2026 F)')."""
    used = re.search(r"\(used (.+)\)$", history)[1]
    entry = next(e for e in history.split("; ") if e.startswith(used + ":"))
    return entry.split(": ", 1)[1].split(" (used")[0]


def report(clash_rows, alternative_rows, unreadable, no_time):
    """Print what most needs a person's eye: the clashes a program student would hit."""
    for term in dict.fromkeys(r[0] for r in clash_rows):
        rows = [r for r in clash_rows if r[0] == term and r[10]]
        counts = {s: sum(1 for r in rows if r[7] == s) for s in SEVERITY}
        print(f"{term}, at the suggested times: "
              + ", ".join(f"{n} {s.lower()}" for s, n in counts.items() if n))
        for r in rows:
            if r[7] in SEVERITY[:2] and r[8] == "same":
                print(f"  {r[7]}: {r[1]} ({r[2]}) × {r[3]} ({r[4]}), {r[5]}")
                print(f"      {r[9]}")
    if alternative_rows:
        print("Each alternative, with everything else at its suggested time:")
        for r in alternative_rows:
            print(f"  {r[0]} {r[1]:<10} {r[2]:<12} {r[3]:<34} {r[5]:>2} clashes, worst: {r[4]}")
    if unreadable:
        print("Suggested times that could not be read (those courses were skipped):")
        for item in unreadable:
            print(f"  {item}")
    if no_time:
        print(f"Named in a program and still offered, but with no in-person time in the last "
              f"{LOOKBACK} falls or winters (not checked):")
        print("  " + ", ".join(no_time))


def about_sheet(sheet, today, args, tt_files, terms, programs):
    about = [
        ("Built", f"{today} by check_timetable.py, in ~/3work/course-scripts"),
        ("Not to be edited", "The three 'Check - ...' sheets are regenerated on every run, and anything "
                             "typed in them is lost. The term sheets and Notes are never touched."),
        ("Programs read", ", ".join(p["name"] for p in programs) + f", from {args.requirements}"),
        ("Timetable exports", f"{len(tt_files)} files, {bc.term_label(terms[0])} to {bc.term_label(terms[-1])}"),
        ("", ""),
        ("Our times", "Suggested time, Alternative time and Tutorial from the term sheets. A tutorial "
                      "is taken to last 50 minutes."),
        ("Other units' times", f"The most recent fall (or winter) in which the course ran in person, "
                               f"looking back at most {LOOKBACK}. 'Its recent times' shows them all, so a "
                               "course that moves every year stands out. Its 2027-28 time is not known."),
        ("Clash", "Every section of one course overlaps every section of the other, so no student "
                  "can dodge it. Lectures and labs/tutorials are both checked."),
        ("Worst case", "Over every program and track: Both required (a student takes both), Required + "
                       "option, Both options, Required + elective (an elective, caught only by a pattern "
                       "like MATH 3***, against a required course), Elective (an elective against an "
                       "option or another elective), "
                       "or Ours only (two of our courses that no program names together)."),
        ("Required or option", "Read from the requirements workbooks: a category whose hours equal the "
                               "hours of its courses is required; one that needs fewer is a choice. "
                               "Courses not offered in the last three years are left out first."),
        ("Years", "'same' when both are first-year, both second-year, or both upper-year. A student "
                  "on the usual path does not take courses from different years together, though "
                  "transfer and part-time students may."),
        ("Programs", "Each program that names both courses, with the category and role of each."),
        ("Suggested, All alternatives", "Whether the clash occurs with every course at its suggested "
                                        "time, and with every alternative taken at once."),
        ("Check - alternatives", "Each course that has an alternative, at each of its times, with "
                                 "everything else at its suggested time."),
        ("", ""),
        ("Limits", "The data shows which courses share a slot, not which students take both."),
        ("", "Prerequisites are not used, so a pair one cannot take together anyway still shows."),
        ("", "Categories that name no courses (first-year science, breadth, Arts electives) are not "
             "read, and nor are the Note columns ('MATH 2513 can count for MATH 2003')."),
        ("", "Substitutes listed beside the usual course (MATH 1053 for MATH 1003, PHYS 1082 for "
             "PHYS 1052) are a short list in the script, SUBSTITUTES; a new one must be added there, "
             "or its category reads as a choice."),
        ("", "No program file covers a student combining Mathematics with Computer Science."),
        ("", "A course with two sections clashes only if both sections do."),
    ]
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 110
    for label, text in about:
        sheet.append([label, text])
    for row in sheet.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")


if __name__ == "__main__":
    main()
