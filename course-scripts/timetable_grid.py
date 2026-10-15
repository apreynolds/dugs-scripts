#!/usr/bin/env -S uv run --script
# /// script
# dependencies = ["openpyxl"]
# ///
# timetable_grid.py — draw one term of the proposed timetable as a week grid.
# Usage:
#   ./timetable_grid.py                     Fall 2027, into timetable-grid-fall-2027.xlsx
#   ./timetable_grid.py -s "Winter 2028"    another term sheet
#   ./timetable_grid.py -o OUT.xlsx
#   ./timetable_grid.py -h | --help
#
#   -w, --workbook   the proposed timetable (default: proposed-timetable-2027-28.xlsx
#                    in timetable-committee). Only read
#   -s, --sheet      the term sheet to draw (default: "Fall 2027")
#   -o, --output     workbook to write (default: timetable-grid-<sheet>.xlsx beside
#                    the proposed timetable). Always written from scratch; a file
#                    already there is first copied, with the date and time in its
#                    name, into the zzzzz folder. Refuses to run while it is open
#
# WHAT THIS IS FOR
# A one-page picture of a term for a department meeting: Monday to Friday
# across the top, each day split into 1000-, 2000-, 3000- and 4000-level
# columns, half-hour rows down the side. Each course sits in the rows its
# suggested time covers, coloured by level, so a crowded slot or a gap stands
# out. The layout is Patrick's, from a sketch he made in the workbook.
#
# WHAT IT SHOWS, AND WHAT IT LEAVES OUT
#   - Suggested times only; alternative times and tutorials are left out.
#   - A course whose sections meet at different times appears at each.
#   - Courses at the same level whose times overlap share one block, which
#     spans all of them; the times are written out when they differ.
#   - MATH 0863, a pre-university course, goes in the 1000-level column.
# The grid is generated: change a time on the term sheet and run this again.
#
# This script borrows its time reading from check_timetable.py and its file
# handling from build_courses.py, which must sit in the same folder.

import argparse
import datetime
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import build_courses as bc
import check_timetable as ct

DAYS = [("M", "MONDAY"), ("T", "TUESDAY"), ("W", "WEDNESDAY"), ("Th", "THURSDAY"), ("F", "FRIDAY")]
LEVELS = [1, 2, 3, 4]
# One colour per level, light enough to print and to read black text on.
COLOURS = {1: "DDEBF7", 2: "E2EFDA", 3: "FFF2CC", 4: "E4DFEC"}  # blue, green, yellow, lilac
FIRST_ROW_TIME = 8 * 60 + 30  # the grid starts at 8:30
STEP = 30  # minutes per row

THIN = Side(style="thin", color="808080")
THICK = Side(style="medium", color="000000")


def time_text(minutes):
    """510 -> '8:30', 780 -> '1:00': the 12-hour clock, as the timetable writes it."""
    hours, mins = divmod(minutes, 60)
    return f"{(hours - 1) % 12 + 1}:{mins:02d}"


def level(key):
    """'MATH 3213' -> 3. MATH 0863 counts as first year."""
    return min(max(int(key[-4]), 1), 4)


def main():
    parser = argparse.ArgumentParser(description="Draw one term of the proposed timetable as a week grid.")
    parser.add_argument("-w", "--workbook", default=ct.DEFAULT_WORKBOOK,
                        help="the proposed timetable (only read)")
    parser.add_argument("-s", "--sheet", default="Fall 2027", help='term sheet to draw (default "Fall 2027")')
    parser.add_argument("-o", "--output", help="workbook to write")
    args = parser.parse_args()

    source = Path(args.workbook).expanduser()
    slug = re.sub(r"\W+", "-", args.sheet.strip().lower())
    output = Path(args.output).expanduser() if args.output else source.with_name(f"timetable-grid-{slug}.xlsx")
    bc.refuse_if_open(output)

    terms = {title: courses for title, _, courses in ct.read_proposal(source)}
    if args.sheet not in terms:
        sys.exit(f"No sheet {args.sheet!r} in {source.name}; it has {', '.join(terms)}.")

    # ---- Where each course goes: (day, level) -> [(first row, last row, label, time)]
    placed = defaultdict(list)
    skipped = []
    latest = FIRST_ROW_TIME
    for key, record in terms[args.sheet].items():
        time = record.get("Suggested time", "")
        options = ct.parse_options(time)
        if not options:
            skipped.append(f"{bc.course_display(key)} {time!r}")
            continue
        label = bc.course_display(key)
        # Several sections at one time read 'MATH 1003 (4 sec)'. Sections at
        # different times are placed separately, so need no count.
        sections = record.get("Sections", "")
        if len(options) == 1 and sections.isdigit() and int(sections) > 1:
            label += f" ({sections} sec)"
        for option in options:
            for day, start, end in option:
                if day not in dict(DAYS):
                    continue
                # A row covers [t, t + 30); the course fills every row it touches.
                first = (start - FIRST_ROW_TIME) // STEP
                last = (end - 1 - FIRST_ROW_TIME) // STEP
                placed[(day, level(key))].append((first, last, label, f"{time_text(start)}-{time_text(end)}"))
                latest = max(latest, end)
    rows = (latest - 1 - FIRST_ROW_TIME) // STEP + 1

    # ---- The sheet ----------------------------------------------------------
    book = Workbook()
    sheet = book.active
    sheet.title = args.sheet
    today = datetime.date.today().isoformat()
    title = sheet.cell(row=1, column=1, value=(
        f"{args.sheet}: suggested times, from {source.name} ({today}). "
        f"Generated by timetable_grid.py; change times on the term sheet and rebuild, not here."))
    title.font = Font(bold=True)
    top = 4  # row of the 8:30 slot; rows 2 and 3 are the day and level headings

    sheet.column_dimensions["A"].width = 7
    for d, (_, day_name) in enumerate(DAYS):
        first_col = 2 + d * len(LEVELS)
        sheet.merge_cells(start_row=2, start_column=first_col, end_row=2, end_column=first_col + len(LEVELS) - 1)
        cell = sheet.cell(row=2, column=first_col, value=day_name)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        for i, lvl in enumerate(LEVELS):
            col = first_col + i
            sheet.column_dimensions[get_column_letter(col)].width = 12
            cell = sheet.cell(row=3, column=col, value=f"{lvl}000-level")
            cell.font = Font(bold=True, size=9)
            cell.alignment = Alignment(horizontal="center")
            cell.fill = PatternFill("solid", start_color=COLOURS[lvl])
    for r in range(rows):
        cell = sheet.cell(row=top + r, column=1, value=time_text(FIRST_ROW_TIME + r * STEP))
        cell.alignment = Alignment(horizontal="right", vertical="top")
        sheet.row_dimensions[top + r].height = 30

    # Faint lines on every cell, so the grid reads as a grid when printed.
    last_col = 1 + len(DAYS) * len(LEVELS)
    for r in range(2, top + rows):
        for c in range(1, last_col + 1):
            # A heavier line on the left of each day.
            left = THICK if c >= 2 and (c - 2) % len(LEVELS) == 0 else THIN
            sheet.cell(row=r, column=c).border = Border(left=left, right=THIN, top=THIN, bottom=THIN)

    # ---- The courses --------------------------------------------------------
    for d, (day, _) in enumerate(DAYS):
        for i, lvl in enumerate(LEVELS):
            col = 2 + d * len(LEVELS) + i
            items = sorted(placed.get((day, lvl), []))
            # Courses whose rows overlap share one block, spanning them all.
            blocks = []
            for item in items:
                # items are sorted by first row, so an overlap can only be with the last block.
                if blocks and item[0] <= max(x[1] for x in blocks[-1]):
                    blocks[-1].append(item)
                else:
                    blocks.append([item])
            for block in blocks:
                first = min(x[0] for x in block)
                last = max(x[1] for x in block)
                same_time = len({x[3] for x in block}) == 1
                lines = []
                for _, _, label, time in block:
                    line = label if same_time else f"{label} {time}"
                    if line not in lines:
                        lines.append(line)
                if last > first:
                    sheet.merge_cells(start_row=top + first, start_column=col, end_row=top + last, end_column=col)
                cell = sheet.cell(row=top + first, column=col, value="\n".join(lines))
                cell.fill = PatternFill("solid", start_color=COLOURS[lvl])
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.font = Font(size=9, bold=True)
                # A dark outline, so two blocks one above the other stay apart.
                for r in range(first, last + 1):
                    target = sheet.cell(row=top + r, column=col)
                    left = THICK if i == 0 else Side(style="thin", color="000000")
                    target.fill = PatternFill("solid", start_color=COLOURS[lvl])
                    target.border = Border(
                        left=left, right=Side(style="thin", color="000000"),
                        top=Side(style="thin", color="000000") if r == first else None,
                        bottom=Side(style="thin", color="000000") if r == last else None,
                    )

    sheet.freeze_panes = sheet.cell(row=top, column=2)
    # Print on one landscape page.
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    backup = bc.save_safely(book, output)
    print(f"Wrote {output}")
    if backup:
        print(f"  The copy it replaced was kept as {backup}")
    if skipped:
        print("  Left out, because the time could not be read: " + ", ".join(skipped))


if __name__ == "__main__":
    main()
