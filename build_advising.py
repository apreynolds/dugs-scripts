# /// script
# dependencies = ["openpyxl"]
# ///
# build_advising.py — build the advising timetable: when the courses our students take usually run.
# Usage:
#   uv run build_advising.py                     rebuild the advising timetable in place
#   uv run build_advising.py -o OUT.xlsx         write a trial copy somewhere else
#   uv run build_advising.py --previous OLD.xlsx -o OUT.xlsx
#   uv run build_advising.py -h | --help
#
#   -o, --output        workbook to write (default: advising-timetable.xlsx in
#                       TEAM-ADVISING/RESOURCES). Always written from scratch.
#                       A file already there is first copied beside it as
#                       NAME-bak.xlsx, replacing the last one. Refuses to
#                       run while the workbook is open in Excel
#   -p, --previous      workbook read ONLY for the typed columns (default: the
#                       output workbook itself, if it exists)
#   -m, --master        the MATH/STAT master list written by build_courses.py
#   -r, --requirements  folder of degree-requirements workbooks
#   -t, --timetable     an Informer "UNB Timetable Data" export, or a directory
#                       of them; repeat the flag for several
#       --years N       how many years the offering notes look back (default 6)
#
# WHAT THIS IS FOR
# One sheet an advisor can open in a meeting: every course that matters to a
# Mathematics/Statistics student, with the slot it usually sits in. It is the
# companion of the master list, and follows the same rule: columns with a
# yellow header are typed by a person and survive every rebuild; everything
# else is regenerated.
#
# WHICH COURSES ARE ON IT
#   1. every course on the MATH/STAT master list (its cells are copied from
#      there, so the two files cannot disagree)
#   2. every other course named in a degree-requirements workbook
#   3. any course typed into the Course column: common electives and the like
# Rows of kind 1 and 2 come back if deleted. A row of kind 3 stays until its
# row is deleted.
#
# ONE ROW PER COURSE
# A course that was renumbered (ESCI 1001 became ESCI 1052), or that is
# cross-listed under a second number (MATH 4443 is also PHYS 4953), is still
# one class in one room. It gets one row, under its current number; the other
# numbers go in that row's "Other numbers" cell. They get no row of their own:
# their sections count as that row's history, and a programme that cites one
# of them is listed in that row's Programs.
#   MATH/STAT rows   the cell is filled from the master list's "Other numbers"
#                    and "Cross-listed" columns, which is where to change it
#   every other row  the cell is typed here
#
# This script borrows its timetable reading from build_courses.py, which must
# sit in the same folder.

import argparse
import datetime
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

import build_courses as bc  # also puts lib/ on sys.path
import paths

# ----------------------------------------------------------------------------
# Defaults, from lib/paths.py
# ----------------------------------------------------------------------------
DEFAULT_OUTPUT = paths.ADVISING_TIMETABLE
DEFAULT_REQUIREMENTS = paths.REQUIREMENTS
DEFAULT_MASTER = paths.MASTER_LIST

# Courses put on the list the first time the workbook is created, because
# Patrick named them as common electives. After that the list is his.
SEED_EXTRAS = ["ECON1013", "ECON1023", "PSYC1013", "PSYC1023", "BIOL1009"]

# What "Listed because" says for a course a person typed in.
HAND = "added by hand"

# A course code as written in a requirements workbook: 'MATH 1003', 'CS1073'.
CODE = re.compile(r"\b([A-Z]{2,4})\s?(\d{4})\b")

# (heading, width, group); "hand" marks the typed columns, as in build_courses.
COLUMNS = [
    ("Course", 11, "hand"), ("Title", 34, "id"), ("CH", 5, "id"),
    ("Offering note", 28, "other"),
    ("Fall timeslot", 38, "fall"), ("Winter timeslot", 38, "winter"),
    ("Summer (in person)", 20, "other"), ("Online", 11, "other"),
    ("Programs", 26, "other"), ("Listed because", 14, "other"),
    ("Status", 10, "other"), ("Curriculum note", 36, "other"),
    ("Other numbers", 26, "hand"), ("Advising note", 44, "hand"),
]


def codes_in(text):
    """Every course number written in a cell, as keys: 'PHYS 1061, PHYS1071' -> two keys."""
    return [subject + number for subject, number in CODE.findall(text or "")]


def master_other_numbers(record):
    """The other numbers a master-list course goes by: (keys, text for the cell).

    The master list keeps these in two columns. "Other numbers" holds numbers
    the registrar files the course under. "Cross-listed" holds its partners in
    other departments, and also its graduate twins (MATH 6443), which are left
    out here: a graduate number means nothing to an undergraduate's timetable.
    """
    keys, parts = [], []
    for subject, number in CODE.findall(record.get("Other numbers", "")):
        keys.append(subject + number)
        parts.append(f"{subject} {number}")
    for subject, number in CODE.findall(record.get("Cross-listed", "")):
        if number[0] >= "6":
            continue  # a graduate twin
        keys.append(subject + number)
        parts.append(f"{subject} {number} (cross-listed)")
    return keys, ", ".join(parts)


def any_course(key):
    """True for anything shaped like a course number, whatever the subject."""
    return re.fullmatch(r"[A-Z]{2,4}\d{4}", key) is not None


def sheet_records(sheet):
    """Read a sheet with a header row into {course: {heading: text}}."""
    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
    found = {}
    if "Course" not in header:
        return found
    for row in rows:
        record = {h: ("" if v is None else str(v).strip()) for h, v in zip(header, row) if h}
        key = bc.course_key(record.get("Course", ""))
        if any_course(key):
            found[key] = record
    return found


def read_requirements(folder):
    """Return {course: [programme names]} for every course a workbook mentions."""
    programs = defaultdict(list)
    for path in sorted(Path(folder).glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue  # Excel's lock file, not a workbook
        book = load_workbook(path, data_only=True)
        if "Requirements" not in book.sheetnames:
            continue
        # 'bsc-math-2020' -> 'bsc-math': short enough to list several in a cell.
        name = re.sub(r"-\d{4}$", "", path.stem)
        for sheet_name in ("Requirements", "Constraints"):
            if sheet_name not in book.sheetnames:
                continue
            for row in book[sheet_name].iter_rows(values_only=True):
                for value in row:
                    for subject, number in CODE.findall(str(value or "")):
                        if name not in programs[subject + number]:
                            programs[subject + number].append(name)
    return programs


def advising_slot(per, season, window, today):
    """The timeslot for one course in fall or winter."""
    slot, _, _ = bc.season_columns(per, season, window, today)
    if slot == "(no set time)":
        # A lab course has no lecture, so every section is one of its lab
        # times. Show those, from the most recent term that has any.
        with_labs = [t for t in sorted(per, key=bc.term_key)
                     if t.endswith(season) and per[t]["inperson"] and per[t]["tutorial_slots"]]
        if with_labs:
            slot = "lab/tut: " + " / ".join(per[with_labs[-1]]["tutorial_slots"])
    # Outside MATH/STAT the extra meeting is as often a lab as a tutorial.
    return slot.replace(" + tut ", " + lab/tut ")


def main():
    parser = argparse.ArgumentParser(
        description="Build the advising timetable from the master list and Informer exports."
    )
    parser.add_argument("-o", "--output", default=DEFAULT_OUTPUT,
                        help="workbook to write (default: the advising timetable itself)")
    parser.add_argument("-p", "--previous",
                        help="workbook to read the typed columns from (default: the output)")
    parser.add_argument("-m", "--master", default=DEFAULT_MASTER,
                        help="the MATH/STAT master list")
    parser.add_argument("-r", "--requirements", default=DEFAULT_REQUIREMENTS,
                        help="folder of degree-requirements workbooks")
    parser.add_argument("-t", "--timetable", action="append",
                        help="timetable export or directory of them (repeatable)")
    parser.add_argument("--years", type=int, default=6,
                        help="years the offering notes look back (default 6)")
    args = parser.parse_args()

    output = Path(args.output).expanduser()
    previous = Path(args.previous).expanduser() if args.previous else output
    bc.refuse_if_open(output)
    today = datetime.date.today().isoformat()

    # ---- What a person typed ------------------------------------------------
    first_build = not previous.exists()
    typed = {}
    if not first_build:
        book = load_workbook(previous, data_only=True)
        typed = sheet_records(book["Timetable"])

    # ---- The three sources of rows ------------------------------------------
    master_book = load_workbook(Path(args.master).expanduser(), data_only=True)
    master = sheet_records(master_book["MAIN"])
    programs = read_requirements(args.requirements)
    # A course on the master list's removed sheet is out of the calendar, so it
    # gets no row here even while a requirements workbook still names it.
    # Those are reported at the end: the workbook is what needs correcting.
    removed = set(sheet_records(master_book["removed"])) if "removed" in master_book.sheetnames else set()
    stale = {key: programs.pop(key) for key in sorted(removed & set(programs))}
    # A row from the previous workbook is kept only if a person added it. The
    # "Listed because" column says so; a row typed in since the last build has
    # that cell blank. Rows listed for another reason must earn their place
    # again, so a course dropped from a requirements workbook drops out here.
    by_hand = {k for k, r in typed.items() if r.get("Listed because", "") in ("", HAND)}

    # ---- Other numbers: old numbers and cross-listed partners ---------------
    # aliases maps a row's course to the other numbers it goes by; other_text
    # is what its "Other numbers" cell shows. See ONE ROW PER COURSE above.
    aliases, other_text = {}, {}
    for key, record in master.items():
        aliases[key], other_text[key] = master_other_numbers(record)
    for key, record in typed.items():
        if key not in master:
            aliases[key] = codes_in(record.get("Other numbers", ""))
            other_text[key] = record.get("Other numbers", "")
    # The other way round: a number -> the row that holds it. A master-list
    # course always keeps its own row, whatever another row claims. If two
    # rows name the same number, both show it, and its sections are counted
    # under whichever row sorts first.
    holder = {}
    for key in sorted(aliases, key=lambda k: (k[:-4], k[-4:])):
        for other in aliases[key]:
            if other not in master and other != key:
                holder.setdefault(other, key)
    # A programme that cites a held number is really citing the row that holds
    # it. This also gives that row its place on the list.
    for other, key in holder.items():
        for name in programs.pop(other, []):
            if name not in programs[key]:
                programs[key].append(name)

    keys = (set(master) | set(programs) | by_hand) - set(holder)
    dropped = [k for k, r in typed.items()
               if k not in keys and (r.get("Other numbers") or r.get("Advising note"))]
    if first_build:
        keys |= set(SEED_EXTRAS)

    # ---- Timetable facts for the courses that are not on the master list ----
    timetable, covered, tt_files = bc.load_timetables(
        args.timetable or bc.DEFAULT_TIMETABLES, subjects=None
    )
    # A row's other numbers count as its own history. Each number is handed to
    # one row only, the one recorded in holder.
    counted = {key: [o for o in others if holder.get(o) == key] for key, others in aliases.items()}
    offerings = bc.build_offerings(timetable, covered, {}, [], counted, {})
    terms = sorted(covered, key=bc.term_key)
    windows = {s: [t for t in terms if t.endswith(s)][-args.years:] for s in ("FA", "WI")}
    years = bc.complete_years(terms, args.years)

    # The latest title and credit hours the registrar has for each number, and
    # which numbers have shared a title (how a renumbering shows up).
    facts = {}
    for (course, term), rows in sorted(timetable.items(), key=lambda kv: bc.term_key(kv[0][1])):
        facts[course] = (rows[0]["Course Title"].strip(), rows[0]["Credit Hours"])
    same_title = defaultdict(set)
    for course, (title, _) in facts.items():
        same_title[(course[:-4], title.lower())].add(course)

    # ---- One row per course --------------------------------------------------
    rows = []
    missing = []
    suggestions = []  # (course, numbers that have carried the same title)
    for key in sorted(keys, key=lambda k: (k[:-4], k[-4:])):
        hand = typed.get(key, {})
        tail = [other_text.get(key, ""), hand.get("Advising note", "")]
        where = ", ".join(sorted(programs.get(key, [])))
        why = "master list" if key in master else "requirements" if key in programs else HAND
        if key in master:
            # Copied, not recomputed: the master list also knows about graduate
            # twins and pre-2017 history, and the two files must agree.
            m = master[key]
            rows.append([bc.course_display(key), m.get("Title", ""), m.get("CH", ""),
                         m.get("Offering note", ""),
                         m.get("Fall timeslot", ""), m.get("Winter timeslot", ""),
                         m.get("Summer (in person)", ""), m.get("Online", ""),
                         where, why,
                         m.get("Status", ""), m.get("Notes", "")] + tail)
            continue
        per = offerings.get(key, {})
        if not per:
            missing.append(key)
            rows.append([bc.course_display(key), "", "", "(not in the timetable since 2017)",
                         "", "", "", "", where, why, "", ""] + tail)
            continue
        # A course that has not yet run under its new number borrows the old one's title.
        title, hours = facts.get(key) or facts.get((aliases.get(key) or [""])[0], ("", ""))
        fall_slot = advising_slot(per, "FA", windows["FA"], today)
        winter_slot = advising_slot(per, "WI", windows["WI"], today)
        # Another number with the same title usually means a renumbering that
        # "Other numbers" does not know about yet. Reported when the run ends.
        others = sorted(same_title.get((key[:-4], title.lower()), set()) - {key} - set(aliases.get(key, [])))
        if others:
            suggestions.append((key, others))
        rows.append([bc.course_display(key), title, str(int(float(hours))) if hours else "",
                     bc.offering_note(per, years), fall_slot, winter_slot,
                     bc.summer_text(per, covered), bc.online_text(per, terms[-1]),
                     where, why, "", ""] + tail)

    # ---- Write ---------------------------------------------------------------
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Timetable"
    bc.write_table(sheet, COLUMNS, rows, freeze="C2")
    bc.shade_by_status(sheet, COLUMNS)

    about = [
        ("Built", f"{today} by build_advising.py"),
        ("Timetable exports", f"{len(tt_files)} files, {bc.term_label(terms[0])} to {bc.term_label(terms[-1])}"),
        ("MATH/STAT rows copied from", str(args.master)),
        ("Programs read from", str(args.requirements)),
        ("", ""),
        ("Yellow headers", "Typed by a person and carried across unchanged on every rebuild."),
        ("Other headers", "Generated. Anything typed there is lost on the next rebuild."),
        ("To add a course", "Type its number in the Course column of a new row and rebuild."),
        ("To remove a course", "Delete its row. A course on the master list or named in a "
                               "degree-requirements workbook comes back on the next rebuild."),
        ("", ""),
        ("Timeslot", "Where the course sat the last time it ran in person in that season. '|' "
                     "separates lectures at different times; '+ lab/tut' lists the lab or tutorial "
                     "options. '-' means never offered in that season. '(new 2026 F)' means the "
                     "slot changed that term, so it is not yet one to rely on."),
        ("Offering note", f"BLANK MEANS EVERY YEAR IN THE SAME TERM, judged over the last "
                          f"{args.years} academic years, fall and winter together. Otherwise it says "
                          "how regular the course is, which term it moves to, or when it last ran."),
        ("Summer, Online", "Summer means in person. Online is never counted as a term offering."),
        ("Programs", "The degree-requirements workbooks that mention the course."),
        ("Listed because", "Why the row is here: it is on the master list, a requirements workbook "
                           "names it, or someone added it by hand. Only rows added by hand are "
                           "carried over from the previous build."),
        ("Other numbers", "Other numbers the same course goes by: one it was renumbered from, or one "
                          "it is cross-listed under. Those numbers get no row of their own; their "
                          "sections and their programmes are counted in this row. For a MATH/STAT "
                          "row the cell is filled from the master list, so change it there; for "
                          "every other row, type it here."),
        ("Status, Curriculum note", "Copied from the MATH/STAT master list, which is where to change them."),
        ("Always check", "This is a guide to what is usual. The online timetable is the authority "
                         "for what is actually scheduled."),
    ]
    info = workbook.create_sheet("About")
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 110
    for label, text in about:
        info.append([label, text])
    for row in info.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    backup = bc.save_safely(workbook, output)
    print(f"Wrote {output}")
    if backup:
        print(f"  The copy it replaced was kept as {backup}")
    outside = len(rows) - len(master)
    print(f"  {len(rows)} courses: {len(master)} from the master list, {outside} others")
    if dropped:
        print("  No longer listed, and their typed notes went with them (see the backup):")
        print("    " + ", ".join(bc.course_display(k) for k in dropped))
    if stale:
        print("  Left out because the master list has them on its removed sheet, though a")
        print("  requirements workbook still names them:")
        for key, names in stale.items():
            print(f"    {bc.course_display(key)}: " + ", ".join(sorted(names)))
    if missing:
        print("  Not found in any timetable export (check the number):")
        print("    " + ", ".join(bc.course_display(k) for k in missing))
    if suggestions:
        print("  Same title under another number (if it is the same course, type the")
        print("  other number in that row's Other numbers):")
        for key, others in suggestions:
            print(f"    {bc.course_display(key)}: " + ", ".join(bc.course_display(o) for o in others))


if __name__ == "__main__":
    main()
