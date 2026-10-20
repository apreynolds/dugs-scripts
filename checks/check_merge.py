# /// script
# dependencies = ["openpyxl", "pdfplumber"]
# ///
# check_merge.py — end-to-end test that a re-run preserves advisor edits.
# Usage: uv run check_merge.py
#
# A regression check, not part of the shipped pipeline. Simulates the yearly workflow:
#   1. generate a tracker from the transcript
#   2. an advisor hand-edits Cat., Ovr., Notes and a gpa1 flag, and renames
#      the student
#   3. the transcript is re-run against that tracker
#   4. every edit must survive, with no duplicated rows

import sys
from pathlib import Path

import openpyxl

# The scripts live one folder up, and their helper modules in lib/.
REPO = Path(__file__).resolve().parent.parent
sys.path[:0] = [str(REPO), str(REPO / "lib")]

import check_support
import tracker_reader


def run(*args):
    return check_support.run_script("transcript_to_tracker.py", *args)


def main():
    WORK, PDF, REQ = check_support.prepare("_merge_test")
    tracker = WORK / "student.xlsx"

    print("STEP 1 — generate")
    run(str(PDF), "-r", str(REQ), "-o", str(tracker), "-t", "Honours")

    print("STEP 2 — simulate advisor edits")
    workbook = openpyxl.load_workbook(str(tracker))
    sheet = workbook["Tracker"]

    # Correct the student's identity (the parser only guesses at it).
    sheet["B1"] = "Sakic, Joe"
    sheet["B2"] = "3213213"

    # Resolve the columns by heading. They were once written as literal
    # numbers, which quietly went one column out of step the moment the Pass
    # column was inserted between Gr. and Ovr.
    col = {sheet.cell(row=5, column=c).value: c
           for c in range(1, sheet.max_column + 1)
           if sheet.cell(row=5, column=c).value}

    # Pick the rows to edit from what is actually in the tracker rather than
    # naming courses. Hardcoded codes tie this check to one transcript, and
    # the fixture is the sort of thing that gets swapped.
    candidates = []
    for row in range(6, 60):
        course = sheet.cell(row=row, column=col["Course"]).value
        grade = sheet.cell(row=row, column=col["Gr."]).value
        term = sheet.cell(row=row, column=col["Term"]).value
        if course and grade and "TRANSFER" not in str(term or "").upper():
            candidates.append((row, str(term), str(course)))

    if len(candidates) < 3:
        raise SystemExit("fixture has too few graded term rows to edit")

    (row_a, term_a, course_a), (row_b, term_b, course_b), \
        (row_c, term_c, course_c) = candidates[:3]

    edits = {}

    # A category + note + gpa1 flag.
    sheet.cell(row=row_a, column=col["Cat. 2020"], value="SciLec")
    sheet.cell(row=row_a, column=col["Notes"],
               value="counts as first-year science")
    sheet.cell(row=row_a, column=col["gpa1"], value=True)
    edits[(term_a, course_a)] = ("SciLec", "",
                                "counts as first-year science", True)

    # An Ovr. that moves a course out of its auto-assigned category — the
    # CS 1003 / CS 1073 case in CLAUDE.md's "looks like a bug but isn't".
    sheet.cell(row=row_b, column=col["Ovr."], value="excl.")
    sheet.cell(row=row_b, column=col["Cat. 2020"], value="FreeElec")
    edits[(term_b, course_b)] = ("FreeElec", "excl.", "", False)

    # A category plus the other GPA flag.
    sheet.cell(row=row_c, column=col["Cat. 2020"], value="Breadth1")
    sheet.cell(row=row_c, column=col["gpa2"], value=True)
    edits[(term_c, course_c)] = ("Breadth1", "", "", False)

    workbook.save(str(tracker))
    print("   edited", len(edits), "rows plus the name/ID:",
          ", ".join(c for _t, c in sorted(edits)))

    before = tracker_reader.read_tracker(tracker)
    n_before = len(before["rows"])

    print("STEP 3 — re-run the same transcript against it")
    output = run(str(PDF), "-r", str(REQ), "-u", str(tracker))
    print("   " + "\n   ".join(
        line for line in output.splitlines() if line.strip()))

    print("STEP 4 — verify")
    after = tracker_reader.read_tracker(tracker)
    by_key = {(r["term"], r["course"]): r for r in after["rows"]}

    failures = []

    if len(after["rows"]) != n_before:
        failures.append("row count changed: " + str(n_before) + " -> " +
                        str(len(after["rows"])))

    # Duplicates are keyed on (term, course), matching the merge key. The same
    # course in two different terms is a RETAKE, not a duplicate; step 5 plants
    # one to prove the two stay distinct.
    keys = [(r["term"], r["course"]) for r in after["rows"] if r["course"]]
    duplicates = sorted({k for k in keys if keys.count(k) > 1})
    if duplicates:
        failures.append("duplicated rows: " +
                        ", ".join(t + " " + c for t, c in duplicates))

    if after["student"]["name"] != "Sakic, Joe":
        failures.append("name not preserved: " +
                        repr(after["student"]["name"]))
    if after["student"]["id"] != "3213213":
        failures.append("id not preserved: " + repr(after["student"]["id"]))
    if after["track"] != "Honours":
        failures.append("track not preserved: " + repr(after["track"]))

    for (term, course), (cat, ovr, notes, gpa1) in edits.items():
        row = by_key.get((term, course))
        if row is None:
            failures.append(course + " vanished")
            continue
        if row.get("cat_2020") != cat:
            failures.append(course + " Cat. = " + repr(row.get("cat_2020")) +
                            ", expected " + repr(cat))
        if row.get("ovr", "") != ovr:
            failures.append(course + " Ovr. = " + repr(row.get("ovr")) +
                            ", expected " + repr(ovr))
        if notes and row.get("notes", "") != notes:
            failures.append(course + " Notes = " + repr(row.get("notes")))
        if gpa1 and not row.get("gpa1"):
            failures.append(course + " gpa1 flag lost")

    if by_key.get((term_c, course_c), {}).get("gpa2") is not True:
        failures.append(course_c + " gpa2 flag lost")

    if failures:
        print()
        print("FAILED:")
        for failure in failures:
            print("  -", failure)
        return 1
    print("   OK — all advisor edits survived, " + str(len(after["rows"])) +
          " rows, no duplicates.")

    # --- the harder cases ---------------------------------------------------

    print()
    print("STEP 5 — delete a whole term, and hand-add a planned course")
    workbook = openpyxl.load_workbook(str(tracker))
    sheet = workbook["Tracker"]

    # Blank out the 2026-1W rows, as though this transcript predated them.
    # NOTE: openpyxl's sheet.cell(row, column, value=None) is a no-op — the
    # value argument is only applied when it is not None — so the cell has to
    # be cleared by assigning to .value directly.
    removed = 0
    for row in range(6, 70):
        if sheet.cell(row=row, column=col["Term"]).value == "2026-1W":
            # Stop at gpa2, the last Courses column. Ranging past it would
            # reach into the requirements block sitting to its right.
            for column in range(1, col["gpa2"] + 1):
                sheet.cell(row=row, column=column).value = None
            removed += 1

    # A course the advisor planned by hand that no transcript will mention.
    first_blank = None
    for row in range(6, 70):
        if not sheet.cell(row=row, column=col["Term"]).value:
            first_blank = row
            break
    sheet.cell(row=first_blank, column=col["Term"], value="2027-1W")
    sheet.cell(row=first_blank, column=col["Course"], value="MATH 4243")
    sheet.cell(row=first_blank, column=col["Title"], value="TOPICS IN ALGEBRA")
    sheet.cell(row=first_blank, column=col["CH"], value=3)
    sheet.cell(row=first_blank, column=col["Cat. 2020"], value="MathElec")
    sheet.cell(row=first_blank, column=col["Notes"],
               value="planned - not yet registered")

    # A RETAKE: the same course code in a different term. The merge key is
    # (term, course), so this must stay a second row rather than collapsing
    # into the original — see invariant 2 in CLAUDE.md. The fixture transcript
    # has no repeated course, so the case has to be planted.
    retake_row = first_blank + 1
    sheet.cell(row=retake_row, column=col["Term"], value="2027-1W")
    sheet.cell(row=retake_row, column=col["Course"], value=course_a)
    sheet.cell(row=retake_row, column=col["Title"], value="RETAKE")
    sheet.cell(row=retake_row, column=col["CH"], value=3)
    sheet.cell(row=retake_row, column=col["Notes"],
               value="retake - planted by the check")

    workbook.save(str(tracker))
    print("   blanked " + str(removed) + " rows from 2026-1W, added a planned "
          "MATH 4243 and a 2027-1W retake of " + course_a)

    print("STEP 6 — re-run")
    output = run(str(PDF), "-r", str(REQ), "-u", str(tracker))
    print("   " + "\n   ".join(
        line for line in output.splitlines() if line.strip()))

    print("STEP 7 — verify the harder cases")
    final = tracker_reader.read_tracker(tracker)
    final_by_key = {(r["term"], r["course"]): r
                    for r in final["rows"]}
    late = []

    if len(final["rows"]) != n_before + 2:
        late.append("expected " + str(n_before + 2) + " rows (the transcript "
                    "plus the planned course and the retake), got " +
                    str(len(final["rows"])))

    # The retake and the original are distinct rows keyed by (term, course).
    same_code = [r for r in final["rows"] if r["course"] == course_a]
    if len(same_code) != 2:
        late.append("the retake of " + course_a + " collapsed into the "
                    "original: " + str(len(same_code)) + " row(s), expected 2")
    elif len({r["term"] for r in same_code}) != 2:
        late.append("both " + course_a + " rows landed in the same term")

    restored = [r for r in final["rows"] if r["term"] == "2026-1W"]
    if len(restored) != removed:
        late.append("2026-1W not fully restored: " + str(len(restored)) +
                    " of " + str(removed))

    planned = final_by_key.get(("2027-1W", "MATH 4243"))
    if planned is None:
        late.append("the hand-added MATH 4243 was deleted")
    else:
        if planned.get("cat_2020") != "MathElec":
            late.append("planned course lost its category")
        if planned.get("notes") != "planned - not yet registered":
            late.append("planned course lost its note")

    # Earlier edits must still be intact after this second merge.
    for (term, course), (cat, _ovr, _notes, _gpa1) in edits.items():
        row = final_by_key.get((term, course))
        if row is None or row.get("cat_2020") != cat:
            late.append(course + " lost its category on the second merge")

    print()
    if late:
        print("FAILED:")
        for failure in late:
            print("  -", failure)
        return 1
    print("PASSED — deleted term re-added, hand-written planned course kept "
          "and reported, earlier edits still intact.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
