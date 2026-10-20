# /// script
# dependencies = ["openpyxl"]
# ///
# check_checkboxes.py — test the native-checkbox post-processing.
# Usage: uv run check_checkboxes.py
#
# A regression check, not part of the shipped pipeline.
#
# checkbox_xlsx.py injects four things into the saved workbook so Excel 365
# renders the gpa1/gpa2 flags as tickable checkboxes. None of it is visible to
# openpyxl and none of it can be evaluated offline, so this checks the
# structure instead:
#
#   1. the part, its content type and its relationship all exist together —
#      any one of them missing is a "repair" prompt;
#   2. <cellXfs count=...> agrees with the number of children;
#   3. EXACTLY the gpa1/gpa2 cells reference a checkbox format, and nothing
#      else does. This is the one that matters. openpyxl dedupes cell formats,
#      so those cells share an <xf> with Term, Course, CH, Gr., Pass and Ovr.;
#      attaching the extension in place instead of cloning would turn six more
#      columns into checkboxes and the XML would look perfectly fine;
#   4. --no-checkboxes produces none of it and restores the TRUE/FALSE
#      dropdowns;
#   5. the flags read back identically either way, and the formulas are
#      untouched — the cell stays an ordinary boolean, which is the whole
#      reason this degrades gracefully on Office 2019 and LibreOffice.

import re
import sys
import zipfile
from pathlib import Path

import openpyxl

# The scripts live one folder up, and their helper modules in lib/.
REPO = Path(__file__).resolve().parent.parent
sys.path[:0] = [str(REPO), str(REPO / "lib")]

import blank_tracker
import check_support
import checkbox_xlsx
import tracker_reader

ROWS = 12

# Cell formats and cells, matched the same attribute-aware way checkbox_xlsx
# matches them — see the comment on its ATTRS for why [^>]* is wrong here.
ATTRS = r'(?:\s+[\w:.-]+="[^"]*")*\s*'
XF_RE = re.compile(r"<xf" + ATTRS + r"(?:/>|>.*?</xf>)", re.S)
CELL_RE = re.compile(r"<c" + ATTRS + r"(?:/>|>)")


def read_part(path, name):
    with zipfile.ZipFile(path) as archive:
        if name not in archive.namelist():
            return None
        return archive.read(name).decode("utf-8")


def part_names(path):
    with zipfile.ZipFile(path) as archive:
        return archive.namelist()


def checkbox_format_indexes(path):
    """Which <cellXfs> entries carry the checkbox extension."""
    styles = read_part(path, "xl/styles.xml")
    block = re.search(r"<cellXfs\b[^>]*>(.*?)</cellXfs>", styles, re.S)
    formats = XF_RE.findall(block.group(1))
    declared = re.search(r'<cellXfs\b[^>]*\bcount="(\d+)"', styles)
    return (formats,
            int(declared.group(1)),
            [i for i, xf in enumerate(formats)
             if checkbox_xlsx.XF_EXT_URI in xf])


def columns_using(path, indexes):
    """Map column letter -> how many of its cells use one of those formats."""
    sheet = read_part(path, "xl/worksheets/sheet1.xml")
    counts = {}
    for tag in CELL_RE.findall(sheet):
        reference = re.search(r'\br="([A-Z]+)\d+"', tag)
        style = re.search(r'\bs="(\d+)"', tag)
        if reference and style and int(style.group(1)) in indexes:
            counts[reference.group(1)] = counts.get(reference.group(1), 0) + 1
    return counts


def formulas(path):
    sheet = read_part(path, "xl/worksheets/sheet1.xml")
    return re.findall(r"<f>(.*?)</f>", sheet, re.S)


def main():
    work = check_support.fresh_dir("_checkbox_test")
    requirements = check_support.sample_requirements(work)

    with_boxes = work / "checkboxes.xlsx"
    without = work / "dropdowns.xlsx"
    blank_tracker.main(["-r", str(requirements), "-o", str(with_boxes),
                        "-t", "Honours", "--rows", str(ROWS)])
    blank_tracker.main(["-r", str(requirements), "-o", str(without),
                        "-t", "Honours", "--rows", str(ROWS),
                        "--no-checkboxes"])

    failures = []

    def fail(message):
        failures.append(message)
        print("  FAIL", message)

    # --- 1. the four pieces are all present, together -----------------------

    print("1. the injected parts")
    names = part_names(with_boxes)
    if checkbox_xlsx.FEATURE_BAG_PART not in names:
        fail("no " + checkbox_xlsx.FEATURE_BAG_PART)
    else:
        print("   OK  part present")

    content_types = read_part(with_boxes, "[Content_Types].xml")
    if checkbox_xlsx.FEATURE_BAG_TYPE not in content_types:
        fail("the part has no content-type override")
    else:
        print("   OK  content type declared")

    rels = read_part(with_boxes, "xl/_rels/workbook.xml.rels")
    if checkbox_xlsx.FEATURE_BAG_REL not in rels:
        fail("the part has no workbook relationship")
    else:
        # A duplicated rId is one of the few ways to earn a repair prompt.
        ids = re.findall(r'\bId="([^"]+)"', rels)
        if len(ids) != len(set(ids)):
            fail("duplicate relationship ids: " + repr(sorted(ids)))
        else:
            print("   OK  relationship present, no duplicate rIds")

    # --- 2. the style table is internally consistent ------------------------

    print()
    print("2. cell formats")
    found, declared, checkbox_indexes = checkbox_format_indexes(with_boxes)
    if declared != len(found):
        fail("<cellXfs count=" + str(declared) + "> but " + str(len(found)) +
             " children — a miscounted clone")
    else:
        print("   OK  count attribute matches (" + str(declared) + ")")

    if not checkbox_indexes:
        fail("no cell format carries the checkbox extension")
    else:
        print("   OK  " + str(len(checkbox_indexes)) +
              " checkbox format(s) appended")

    # extLst must be the LAST child of <xf>, after alignment and protection.
    for index in checkbox_indexes:
        body = found[index]
        for element in ("</alignment>", "</protection>"):
            if element in body and body.index(element) > body.index("<extLst>"):
                fail("extLst is not the last child of <xf> " + str(index))

    # --- 3. the clone-not-mutate property -----------------------------------

    print()
    print("3. only the gpa flags became checkboxes")
    sheet = openpyxl.load_workbook(str(with_boxes))["Tracker"]
    headings = {sheet.cell(row=5, column=c).value: c
                for c in range(1, sheet.max_column + 1)
                if sheet.cell(row=5, column=c).value}
    wanted = {openpyxl.utils.get_column_letter(headings[name])
              for name in ("gpa1", "gpa2")}

    used = columns_using(with_boxes, set(checkbox_indexes))
    if set(used) != wanted:
        fail("checkbox formats are used by columns " + repr(sorted(used)) +
             ", expected exactly " + repr(sorted(wanted)) +
             " — the shared <xf> was mutated instead of cloned")
    else:
        print("   OK  only " + ", ".join(sorted(wanted)) + " (" +
              ", ".join(str(used[c]) + " cells in " + c
                        for c in sorted(used)) + ")")

    # The columns that share the ORIGINAL format are the ones at risk.
    neighbours = {openpyxl.utils.get_column_letter(headings[name])
                  for name in ("Term", "Course", "CH", "Gr.", "Pass", "Ovr.")
                  if name in headings}
    caught = neighbours & set(used)
    if caught:
        fail("these columns were dragged along: " + repr(sorted(caught)))
    else:
        print("   OK  " + ", ".join(sorted(neighbours)) + " left alone")

    # --- 4. the opt-out really opts out -------------------------------------

    print()
    print("4. --no-checkboxes")
    if checkbox_xlsx.FEATURE_BAG_PART in part_names(without):
        fail("--no-checkboxes still injected the part")
    elif checkbox_xlsx.FEATURE_BAG_TYPE in read_part(
            without, "[Content_Types].xml"):
        fail("--no-checkboxes still declared the content type")
    else:
        print("   OK  none of the parts are present")

    plain = openpyxl.load_workbook(str(without))["Tracker"]
    gpa_letters = {openpyxl.utils.get_column_letter(headings[name])
                   for name in ("gpa1", "gpa2")}
    covered = set()
    for validation in plain.data_validations.dataValidation:
        if validation.formula1 == '"TRUE,FALSE"':
            for cell_range in validation.sqref.ranges:
                for column in range(cell_range.min_col, cell_range.max_col + 1):
                    covered.add(openpyxl.utils.get_column_letter(column))
    if not gpa_letters <= covered:
        fail("the TRUE/FALSE dropdown is missing from " +
             repr(sorted(gpa_letters - covered)))
    else:
        print("   OK  TRUE/FALSE dropdowns restored on " +
              ", ".join(sorted(gpa_letters)))

    # And the checkbox build must NOT carry that dropdown — an arrow beside a
    # checkbox reads as a second, different control.
    boxed_covered = set()
    for validation in sheet.parent["Tracker"].data_validations.dataValidation:
        if validation.formula1 == '"TRUE,FALSE"':
            for cell_range in validation.sqref.ranges:
                for column in range(cell_range.min_col, cell_range.max_col + 1):
                    boxed_covered.add(
                        openpyxl.utils.get_column_letter(column))
    if gpa_letters & boxed_covered:
        fail("the checkbox build still has a dropdown on " +
             repr(sorted(gpa_letters & boxed_covered)))
    else:
        print("   OK  no dropdown on the checkbox flags")

    # --- 5. nothing else changed --------------------------------------------

    print()
    print("5. the workbook is otherwise identical")
    if formulas(with_boxes) != formulas(without):
        fail("the formulas differ between the two builds")
    else:
        print("   OK  every formula identical (" +
              str(len(formulas(with_boxes))) + ")")

    boxed_rows = tracker_reader.read_tracker(with_boxes)
    plain_rows = tracker_reader.read_tracker(without)
    if boxed_rows != plain_rows:
        fail("tracker_reader sees the two builds differently")
    else:
        print("   OK  tracker_reader reads both the same")

    values = {sheet.cell(row=r, column=headings["gpa1"]).value
              for r in range(6, 6 + ROWS)}
    if values != {False}:
        fail("gpa1 cells are not plain booleans: " + repr(values))
    else:
        print("   OK  the cells are still ordinary booleans")

    # --- 6. the documented limitation ---------------------------------------
    #
    # Pinned rather than merely written down: openpyxl rebuilds the archive
    # from its object model, so loading a checkbox workbook and saving it
    # DROPS the checkboxes. That is fine for this pipeline — a tracker is
    # always generated from scratch, never read-modify-written (invariant 1) —
    # but anyone who reaches for openpyxl to tweak a finished tracker should
    # find this stated rather than discover it.

    print()
    print("6. openpyxl round-trip drops them (expected)")
    round_tripped = work / "round-tripped.xlsx"
    openpyxl.load_workbook(str(with_boxes)).save(str(round_tripped))
    if checkbox_xlsx.FEATURE_BAG_PART in part_names(round_tripped):
        print("   note: openpyxl now PRESERVES the part — the warning in "
              "CLAUDE.md can be softened")
    else:
        print("   OK  confirmed: re-saving with openpyxl strips them, so "
              "never read-modify-write a tracker")

    print()
    if failures:
        print("FAILED —", len(failures), "problem(s) above")
        return 1
    print("PASSED — checkboxes injected on the gpa flags only, and only when "
          "asked for.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
